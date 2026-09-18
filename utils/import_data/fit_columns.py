#!/usr/bin/env python3
"""按内容自动设置列宽，并冻结首行与左侧标识列。

用法：
    python3 fit_columns.py --xlsx A.xlsx --sheet 去重后总表
    python3 fit_columns.py --xlsx A.xlsx --sheet 去重后总表 --freeze-cols 6

**中日韩字符按两个字符宽计**。Excel 的列宽单位约等于「一个半角字符」，
直接拿 len() 算，中文列会窄掉近一半 —— 合并表里展品名称、简介全是中英混排，
不处理的话看起来就是挤成一团。

**按分位数而不是最大值算**（--pct，默认 95）：典型值与极值差得很远 ——
`展厅` 列中位数 16、p90 却是 51，最长 84。按最大值算，一个长值就把整列撑爆，
其余几千行跟着陪绑。取 p95 让绝大多数行舒服，个别超长的截断显示。

**长文本列封顶**（--max-width）：展品简介动辄几百字，按内容撑开会把整张表
推到屏幕外。封顶后靠单元格自身滚动查看。

**冻结列别贪多**：冻 6 列时冻结区宽 153 字符，比多数屏幕还宽，横向滚动就失效了。
默认冻到「序号」为止。
"""
from __future__ import annotations

import argparse
import pathlib
import shutil
import sys
import unicodedata

import openpyxl
from openpyxl.utils import get_column_letter


def wide(ch: str) -> int:
    """East Asian Wide / Fullwidth 算两格。"""
    return 2 if unicodedata.east_asian_width(ch) in ("W", "F") else 1


def disp_len(v) -> int:
    if v is None:
        return 0
    s = str(v)
    # 多行取最长的一行 —— 按总长算会把换行文本撑得离谱
    return max((sum(wide(c) for c in line) for line in s.splitlines()), default=0)


def main() -> None:
    ap = argparse.ArgumentParser()
    ap.add_argument("--xlsx", required=True)
    ap.add_argument("--sheet", required=True)
    ap.add_argument("--min-width", type=int, default=7)
    ap.add_argument("--max-width", type=int, default=40)
    ap.add_argument("--pct", type=int, default=95,
                    help="按第几百分位的内容长度定宽（默认 95）")
    ap.add_argument("--sample", type=int, default=800, help="按前 N 行取样算宽度")
    ap.add_argument("--freeze-cols", type=int, default=4,
                    help="冻结左侧几列（默认 4，冻到「序号」）。0 表示只冻结首行")
    args = ap.parse_args()

    path = pathlib.Path(args.xlsx)
    if not path.exists():
        sys.exit(f"文件不存在：{path}")
    wb = openpyxl.load_workbook(path)
    if args.sheet not in wb.sheetnames:
        sys.exit(f"没有 sheet「{args.sheet}」，现有：{wb.sheetnames}")
    ws = wb[args.sheet]

    lens: dict[int, list[int]] = {}
    for r, row in enumerate(ws.iter_rows(), start=1):
        if r > args.sample + 1:
            break
        for cell in row:
            n = disp_len(cell.value)
            if n:
                lens.setdefault(cell.column, []).append(n)

    widths: dict[int, int] = {}
    for col, xs in lens.items():
        xs.sort()
        k = max(0, min(len(xs) - 1, int(len(xs) * args.pct / 100) - 1))
        # 表头本身要放得下，否则列名被截断，翻到第 80 列时根本不知道在看什么
        widths[col] = max(xs[k], disp_len(ws.cell(1, col).value))

    for col, n in widths.items():
        w = max(args.min_width, min(args.max_width, n + 2))
        ws.column_dimensions[get_column_letter(col)].width = w

    ws.freeze_panes = f"{get_column_letter(args.freeze_cols + 1)}2" if args.freeze_cols else "A2"

    bak = path.with_suffix(".xlsx.bak_width")
    if not bak.exists():
        shutil.copy2(path, bak)
        print(f"改动前已备份到 {bak.name}")
    wb.save(path)
    hdr = [c.value for c in ws[1]]
    print(f"已设置 {len(widths)} 列的列宽（{args.min_width}–{args.max_width}），"
          f"冻结 {args.freeze_cols} 列 + 首行；按 p{args.pct} 定宽")
    print("前 10 列宽度：")
    for i in range(1, min(11, len(hdr) + 1)):
        print(f"  {get_column_letter(i):<3} {str(hdr[i-1])[:18]:<20} "
              f"{ws.column_dimensions[get_column_letter(i)].width}")


if __name__ == "__main__":
    main()
