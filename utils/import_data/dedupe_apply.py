#!/usr/bin/env python3
"""把去重结论写回工作簿：加三列标记 + 另建「去重后总表_v2」。

用法：
    python3 dedupe_apply.py --xlsx "exports/展品_波士顿美术馆_展厅检索.xlsx" --dry-run
    python3 dedupe_apply.py --xlsx "..."

**绝不就地删行。** `copy_sheet.py` 的文件头已记：`fill_onview_gemini.py` /
`gallery_grounded.py` 都按 xlsx 行号定位（`序号` 在合并表里不唯一），
行序一变，A 件的答案会写进 B 件的行**且不报错**。所以原 sheet 只在末尾加列，
删行版本另开一张 sheet，回滚 = 删掉三列与新 sheet。

三列的取值刻意把「没查过」和「查过且独一」分开：

| 去重状态 | 含义 |
|---|---|
| `主行` | 一组重复里保留的那条 |
| `并入(组ID)` | 被并入主行的那条，v2 里不出现 |
| `已复核·独一` | 在「同作者同时期」组里送过模型，判为独一件 |
| `待人工复核` | 模型判同一件，但被确定性闸挡下（簇内年代冲突 / 类型不符 / 佚名且原清单一侧无任何锚点） |
| `未复核·同组全为Wikidata条目` | 同组其它件都是不同的 Wikidata 实体，默认不送模型 |
| `未复核·无同组` | 该作者该时期只有这一件 —— **没有可比对象**，与「已复核·独一」是两回事 |

上一轮把 151 件零候选默认当成「独有」，正是因为这两种在数据里长得一样。

**主行怎么选**（用户 2026-09-16 定）：**按非空字段数**，多的当主行。
不按来源固定优先级 —— ext·官网 常有中文题名+展厅+材质，ext·wikidata 独有 QID+馆藏号，
两边各有所长，固定优先级必然丢信息。落败方的独有字段进 `X_补充N` 列
（沿用表里现有的 52 个补充列同一套机制）。
"""
from __future__ import annotations

import argparse
import collections
import csv
import json
import pathlib
import shutil
import itertools
import sys

import openpyxl

import dedupe_lib as L
from gallery_text import precision

HERE = pathlib.Path(__file__).parent
HARD = HERE / "dedupe_hard.json"
CLUSTERS = HERE / "dedupe_clusters.json"
REPORT = HERE / "dedupe_report.csv"
SHEET = "去重后总表"
SHEET_V2 = "去重后总表_v2"
SHEET_LIST = "重复展品清单"
SHEET_DETAIL = "重复展品明细"
NEW_COLS = ("去重组ID", "去重状态", "去重依据")


def load_rows(path: pathlib.Path):
    wb = openpyxl.load_workbook(path, read_only=True, data_only=True)
    ws = wb[SHEET]
    it = ws.iter_rows(values_only=True)
    hdr = list(next(it))
    rows = {}
    for n, r in enumerate(it, start=2):
        d = dict(zip(hdr, r))
        rows[(d.get("来源表"), d.get("来源行"))] = {"row": n, "data": d}
    return hdr, rows


def row_facts(path: pathlib.Path) -> dict[tuple, dict]:
    """(来源表,来源行) → 统一事实（年代、子集、馆藏号）。与分组用的是同一套事实。"""
    import dedupe_group as G
    items, _ = G.load_items(str(path))
    return {i["key"]: i for i in items}


def nonempty(d: dict) -> int:
    return sum(1 for v in d.values() if v not in (None, ""))


def main() -> None:
    ap = argparse.ArgumentParser()
    ap.add_argument("--xlsx", required=True)
    ap.add_argument("--dry-run", action="store_true")
    args = ap.parse_args()

    path = pathlib.Path(args.xlsx)
    hdr, rows = load_rows(path)
    n_before = len(rows)
    print(f"读入 {n_before} 行 × {len(hdr)} 列")

    # ---------------- 收集「判为同一件」的边
    edges, why = [], {}
    for h in json.loads(HARD.read_text(encoding="utf-8")):
        if h.get("auto"):
            a, b = tuple(h["a"]), tuple(h["b"])
            edges.append((a, b))
            why[(a, b)] = h["rule"]
    n_hard = len(edges)
    # 分组结论（用户 2026-09-17 定的策略：作者+时期分组，haiku 组内判重）。
    # **不再读 dedupe_pairs.json** —— 那是旧策略（逐对、opus-5）的产物，
    # 与现在的候选集合对不上，混用会把两套判据的结论搅在一起。
    if not CLUSTERS.exists():
        sys.exit(f"缺 {CLUSTERS.name} —— 先跑 dedupe_group.py")
    cdata = json.loads(CLUSTERS.read_text(encoding="utf-8"))
    by_row = {v["row"]: k for k, v in rows.items()}
    facts = row_facts(path)
    spans = {k: v['span'] for k, v in facts.items()}
    sent = {by_row[r] for r in cdata["sent_rows"] if r in by_row}
    skipped = {by_row[r] for r in cdata.get("skipped_ext_rows", []) if r in by_row}

    rejected, same_model = [], 0
    for c in cdata["clusters"]:
        keys = [by_row[r] for r in c["rows"] if r in by_row]
        # ⚠ 确定性闸：簇内任两件年代冲突，整簇作废。
        # 「重复只可能发生在同一时期」由这里保证，不靠模型自觉。
        bad = [(a, b) for a, b in itertools.combinations(keys, 2)
               if L.year_relation(spans.get(a), spans.get(b)) == "conflict"]
        if bad:
            rejected.append((keys, c, "簇内年代冲突"))
            continue
        # ⚠ 类型闸：一边平面（画、版画）、一边立体（雕塑、器物），不可能是同一件。
        types = {L.coarse_type(rows[k]["data"]) for k in keys} - {None}
        if len(types) > 1:
            rejected.append((keys, c, "类型不符（一件平面、一件立体）"))
            continue
        # ⚠ 弱证据闸（佚名簇）：原清单的行只是一个描述性的中文短名，
        # 若它**既无年代又无馆藏号**，就没有任何能指向馆里具体某一件的锚点 ——
        # 「圣塞巴斯蒂安」「三博士朝拜」「黄铜星盘」，馆里同题材的往往不止一件。
        # 具名作者的簇不受此限：作者+时期分组本身已把范围收得很窄。
        # 实测 haiku 对「黄铜星盘 ↔ An astrolabe…」给的理由就是「同一类文物」—— 同类不是同件。
        if c["akey"] == "佚名":
            weak = [k for k in keys if facts[k]["sub"] == "原清单"
                    and not facts[k]["span"] and not any(facts[k]["acc"].values())]
            if weak:
                rejected.append((keys, c, "佚名且原清单一侧无年代无馆藏号"))
                continue
        for b in keys[1:]:
            edges.append((keys[0], b))
            # setdefault：同一对已有硬证据（馆藏号/组画）时保留硬证据的说明 ——
            # 否则 haiku 的说明会覆盖它，清单里「强」证据被误标成「中」。
            k1, k2 = (keys[0], b), (b, keys[0])
            if k1 not in why and k2 not in why:
                why[k1] = f"作者+时期同组·haiku 判同一件（{c['akey']}）：{c['reason'][:50]}"
        same_model += 1
    print(f"判为同一件：硬证据 {n_hard} 对 + 组内判重 {same_model} 簇"
          f"（模型 {cdata.get('model')}）")
    print(f"被闸挡下、降级为待人工复核的簇：{len(rejected)}")
    for keys, c, why_bad in rejected:
        print(f"    [{why_bad}] {' ↔ '.join(str(rows[k]['data'].get('展品名称') or '')[:18] for k in keys)}")

    # ---------------- 并查集成组
    parent: dict[tuple, tuple] = {}

    def find(x):
        parent.setdefault(x, x)
        while parent[x] != x:
            parent[x] = parent[parent[x]]
            x = parent[x]
        return x

    for a, b in edges:
        ra, rb = find(a), find(b)
        if ra != rb:
            parent[rb] = ra
    groups: dict[tuple, list] = collections.defaultdict(list)
    for k in list(parent):
        groups[find(k)].append(k)
    groups = {r: sorted(set(m)) for r, m in groups.items() if len(set(m)) > 1}
    print(f"成组 {len(groups)} 组，涉及 {sum(len(m) for m in groups.values())} 行")

    # ---------------- 定主行：非空字段数最多的那条
    status: dict[tuple, tuple[str, str, str]] = {}
    merged_into: dict[tuple, tuple] = {}
    for n, (root, members) in enumerate(sorted(groups.items()), start=1):
        gid = f"G{n:04d}"
        members = [m for m in members if m in rows]
        if len(members) < 2:
            continue
        main_key = max(members, key=lambda k: nonempty(rows[k]["data"]))
        cand = [v for (a, b), v in why.items() if a in members and b in members]
        basis = next((v for v in cand if v.startswith(("跨管线", "组画"))), cand[0] if cand else "")
        status[main_key] = (gid, "主行", basis)
        for m in members:
            if m == main_key:
                continue
            status[m] = (gid, f"并入({gid})", basis)
            merged_into[m] = main_key

    # ---------------- 其余行：把「没查过」与「查过且独一」分开
    rej_why = {k: w for keys, _, w in rejected for k in keys}
    for k in rows:
        if k in status:
            continue
        if k in rej_why:
            status[k] = ("", "待人工复核", f"模型判同一件，但被闸挡下：{rej_why[k]}")
        elif k in sent:
            status[k] = ("", "已复核·独一", "同作者同时期组内，模型判为独一件")
        elif k in skipped:
            status[k] = ("", "未复核·同组全为Wikidata条目",
                         "同作者同时期的其它件都是不同的 Wikidata 实体，未送模型")
        else:
            status[k] = ("", "未复核·无同组", "该作者该时期只有这一件（或佚名且无共同罕见题名词）")
    dist = collections.Counter(v[1].split("(")[0] for v in status.values())
    print(f"状态分布：{dict(dist)}")

    # ---------------- 报告
    with REPORT.open("w", encoding="utf-8-sig", newline="") as f:
        w = csv.writer(f)
        w.writerow(["去重组ID", "去重状态", "去重依据", "来源表", "来源行",
                    "Tier", "展品名称", "馆藏号", "官方页面"])
        for k, (gid, st, basis) in sorted(status.items(), key=lambda kv: (kv[1][0] == "", kv[0])):
            if not st.startswith(("主行", "并入", "待人工复核")):
                continue
            d = rows[k]["data"]
            w.writerow([gid, st, basis, k[0], k[1], d.get("Tier"),
                        str(d.get("展品名称") or "")[:90], d.get("馆藏号"), d.get("官方页面")])
    print(f"已写出 {REPORT.name}")

    if args.dry_run:
        print("\n--dry-run：不写工作簿。成组样例 —")
        for n, (root, members) in enumerate(sorted(groups.items())[:6], start=1):
            print(f"  组 {n}：")
            for m in members:
                if m in rows:
                    mark = "主" if status.get(m, ("", ""))[1] == "主行" else "并"
                    print(f"    [{mark}] {m[0]}表第{m[1]}行  "
                          f"{str(rows[m]['data'].get('展品名称') or '(空)')[:52]}")
        return

    # ---------------- 写盘：原 sheet 末尾加三列 + 新建 v2
    wb = openpyxl.load_workbook(path)
    ws = wb[SHEET]
    head = [c.value for c in ws[1]]
    base = len(head)
    # 幂等：列已存在就就地覆盖，不重复追加
    col_at = {}
    for name in NEW_COLS:
        if name in head:
            col_at[name] = head.index(name) + 1
        else:
            base += 1
            ws.cell(1, base).value = name
            col_at[name] = base
    ia, ib = head.index("来源表") + 1, head.index("来源行") + 1
    for r in range(2, ws.max_row + 1):
        k = (ws.cell(r, ia).value, ws.cell(r, ib).value)
        gid, st, basis = status.get(k, ("", "", ""))
        for name, v in zip(NEW_COLS, (gid, st, basis)):
            ws.cell(r, col_at[name]).value = v or None

    if SHEET_V2 in wb.sheetnames:
        del wb[SHEET_V2]
    v2 = wb.create_sheet(SHEET_V2)
    head2 = [c.value for c in ws[1]]
    i_st = head2.index("去重状态")
    i_src = head2.index("重复合并来源") if "重复合并来源" in head2 else None
    # 主行 → 被并入行的原始值，用于把落败方的独有字段搬进补充列
    raw = {}
    for r in range(2, ws.max_row + 1):
        raw[(ws.cell(r, ia).value, ws.cell(r, ib).value)] = \
            [ws.cell(r, c).value for c in range(1, ws.max_column + 1)]
    # 身份列不搬：被并入行的来源与序号已经记在「重复合并来源」里，
    # 再往补充列里塞一遍只会把真正有价值的位置挤掉。
    SKIP_CARRY = {"来源表", "来源行", "序号", "博物馆", *NEW_COLS}

    # `<列名>_补充<N>` 是表里现有的机制（52 个补充列），沿用；**不够就扩**。
    # 早先不扩，217 个值被静默丢弃，其中 29 个是「展品名称」——
    # 那正是另一语种的题名，是这份数据双语可用的关键，丢了就白合并了。
    slots: dict[str, list[int]] = collections.defaultdict(list)
    for j, h in enumerate(head2):
        if h and "_补充" in str(h):
            slots[str(h).split("_补充")[0]].append(j)

    need = collections.Counter()
    for m, tgt in merged_into.items():
        md, od = rows.get(tgt, {}).get("data", {}), rows.get(m, {}).get("data", {})
        for h in head2:
            if not h or "_补充" in str(h) or h in SKIP_CARRY:
                continue
            ov, mv = od.get(h), md.get(h)
            if ov not in (None, "") and mv not in (None, "") and ov != mv:
                need[str(h)] += 1
    extra = 0
    for c, n in need.items():
        have = len(slots.get(c, []))
        for i in range(have, n):
            head2.append(f"{c}_补充{i + 1}")
            slots[c].append(len(head2) - 1)
            extra += 1
    if extra:
        print(f"为放下落败方的独有字段，新增 {extra} 个补充列")

    v2.append(head2)
    kept = 0
    carried = collections.Counter()
    for r in range(2, ws.max_row + 1):
        vals = [ws.cell(r, c).value for c in range(1, ws.max_column + 1)]
        vals += [None] * (len(head2) - len(vals))        # 补齐新增的补充列
        if str(vals[i_st] or "").startswith("并入"):
            continue
        k = (vals[ia - 1], vals[ib - 1])
        same = [m for m, tgt in merged_into.items() if tgt == k]
        for m in same:
            other = raw.get(m)
            if not other:
                continue
            for j, h in enumerate(head2):
                if not h or "_补充" in str(h) or h in SKIP_CARRY:
                    continue
                ov = other[j]
                if ov in (None, "") or ov == vals[j]:
                    continue
                if vals[j] in (None, ""):
                    vals[j] = ov                      # 主行为空，直接补上
                    carried["填空"] += 1
                else:
                    # 展厅列有粒度可比：**更具体的当主值**，主行原值退到补充列。
                    # 不这么做的话，G0025 的主行会留着 `Europe`，
                    # 而被并入那行更准的 `Art of Europe 欧洲艺术（英国绘画）` 被挤进补充列 ——
                    # 合并本该是取两边之长，不是让主行的身份决定取值。
                    if str(h) == "展厅" and precision(ov) > precision(vals[j]):
                        vals[j], ov = ov, vals[j]
                        carried["展厅升级"] += 1
                    free = next((s for s in slots.get(str(h), [])
                                 if vals[s] in (None, "")), None)
                    if free is not None:
                        vals[free] = ov
                        carried["进补充列"] += 1
                    else:
                        carried["无补充列可放·丢弃"] += 1
        # 被并入行的 (来源表,来源行) 记进「重复合并来源」，沿用现有 61 行的格式
        if i_src is not None and same:
            old = str(vals[i_src] or "").strip()
            vals[i_src] = (old + "；" if old else "") + \
                "；".join(f"{m[0]}表第{m[1]}行" for m in same)
        v2.append(vals)
        kept += 1
    print(f"字段搬运：{dict(carried)}")
    v2.freeze_panes = "A2"

    # ---------------- 「重复展品清单」：给人看的那一份
    # 每组一段，主行在前；证据强度分三档，让人知道哪些该回头核一眼。
    # 被闸挡下的「待人工复核」附在最后 —— 它们是模型说同一件、但证据不够的。
    if SHEET_LIST in wb.sheetnames:
        del wb[SHEET_LIST]
    ls = wb.create_sheet(SHEET_LIST, 0)
    ls.append(["组ID", "判定方式", "证据强度", "保留/并入", "展品名称", "作者", "年代",
               "馆藏号", "来源表", "来源行", "Tier", "判定依据"])

    def fmt_span(sp):
        if not sp:
            return ""
        lo, hi, ap = sp
        y = lambda v: f"公元前{-v}" if v < 0 else str(v)
        return ("约" if ap else "") + y(lo) + ("" if hi == lo else f"–{y(hi)}")

    def strength(members, basis):
        if basis.startswith(("跨管线", "组画")):
            return "强"
        if all(not spans.get(m) for m in members):
            return "弱·两侧无年代，建议人工确认"
        return "中"

    gid_members = collections.defaultdict(list)
    for k, (gid, st, basis) in status.items():
        if gid:
            gid_members[gid].append(k)
    for gid in sorted(gid_members):
        members = sorted(gid_members[gid], key=lambda k: status[k][1] != "主行")
        basis = status[members[0]][2]
        how = ("硬证据·跨管线馆藏号" if basis.startswith("跨管线") else
               "硬证据·组画分件" if basis.startswith("组画") else "haiku·作者+时期同组")
        for m in members:
            d, fx = rows[m]["data"], facts.get(m, {})
            acc = sorted(set().union(*fx.get("acc", {}).values())) if fx else []
            ls.append([gid, how, strength(members, basis),
                       "保留" if status[m][1] == "主行" else "并入",
                       str(d.get("展品名称") or "（名称为空）"),
                       fx.get("artist_txt") or "", fmt_span(spans.get(m)),
                       ", ".join(acc[:2]), m[0], m[1], d.get("Tier"),
                       basis.split("：", 1)[-1] if m == members[0] else ""])
        ls.append([])
    pend = [k for k, v in status.items() if v[1] == "待人工复核"]
    if pend:
        ls.append(["待人工复核 —— 模型判为同一件，但被确定性闸挡下，未合并"])
        for k in sorted(pend):
            d = rows[k]["data"]
            ls.append(["", "", "", "待复核", str(d.get("展品名称") or ""), "",
                       fmt_span(spans.get(k)), "", k[0], k[1], d.get("Tier"),
                       status[k][2]])
    for col, w in zip("ABCDEFGHIJKL", (8, 20, 26, 9, 60, 22, 14, 16, 7, 7, 6, 60)):
        ls.column_dimensions[col].width = w
    ls.freeze_panes = "A2"
    print(f"已生成「{SHEET_LIST}」：{len(gid_members)} 组 + 待复核 {len(pend)} 行（放在第一张）")

    # ---------------- 「重复展品明细」：参与去重的每一条记录，原表整行
    # 清单只有摘要列；这张给要逐字段对照的人看。被并入的行在 v2 里已经没有了，
    # 这里是它们唯一还能整行看到的地方（原 sheet 里也有，但散在 4537 行中间）。
    if SHEET_DETAIL in wb.sheetnames:
        del wb[SHEET_DETAIL]
    dt = wb.create_sheet(SHEET_DETAIL, 1)
    hd = [c.value for c in ws[1]]
    lead = ["组ID", "保留/并入", "判定方式", "证据强度"]
    dt.append(lead + [h for h in hd if h not in NEW_COLS] + ["去重依据"])
    keep_idx = [j for j, h in enumerate(hd) if h not in NEW_COLS]
    i_basis = hd.index("去重依据")
    n_detail = 0
    for gid in sorted(gid_members):
        members = sorted(gid_members[gid], key=lambda k: status[k][1] != "主行")
        basis = status[members[0]][2]
        how = ("硬证据·跨管线馆藏号" if basis.startswith("跨管线") else
               "硬证据·组画分件" if basis.startswith("组画") else "haiku·作者+时期同组")
        for m in members:
            vals = raw[m]
            dt.append([gid, "保留" if status[m][1] == "主行" else "并入", how,
                       strength(members, basis)]
                      + [vals[j] for j in keep_idx] + [vals[i_basis]])
            n_detail += 1
        dt.append([])
    if pend:
        dt.append(["待人工复核 —— 模型判为同一件，但被确定性闸挡下，未合并"])
        for k in sorted(pend):
            vals = raw[k]
            dt.append(["", "待复核", "", ""] + [vals[j] for j in keep_idx] + [vals[i_basis]])
            n_detail += 1
    widths = {"组ID": 8, "保留/并入": 9, "判定方式": 20, "证据强度": 16,
              "展品名称": 56, "展品简介": 40, "去重依据": 50}
    from openpyxl.utils import get_column_letter
    for j, h in enumerate(lead + [h for h in hd if h not in NEW_COLS] + ["去重依据"], start=1):
        dt.column_dimensions[get_column_letter(j)].width = widths.get(h, 14)
    dt.freeze_panes = "E2"
    print(f"已生成「{SHEET_DETAIL}」：{n_detail} 条记录（原表整行，{len(keep_idx)} 列）")

    # ---------------- 写盘前硬断言：静默降级比崩掉危险得多
    n_merged = sum(1 for v in status.values() if v[1].startswith("并入"))
    if ws.max_row - 1 != n_before:
        sys.exit(f"原 sheet 行数变了（{ws.max_row - 1} ≠ {n_before}）—— 未保存")
    if kept != n_before - n_merged:
        sys.exit(f"v2 行数对不上（{kept} ≠ {n_before} − {n_merged}）—— 未保存")

    bak = path.with_suffix(".xlsx.bak_dedupe")
    if not bak.exists():
        shutil.copy2(path, bak)
        print(f"改动前已备份到 {bak.name}")
    wb.save(path)
    print(f"已写回 {path.name}：原 sheet {n_before} 行 × {ws.max_column} 列（+3），"
          f"「{SHEET_V2}」{kept} 行（并入 {n_merged} 行）")


if __name__ == "__main__":
    main()
