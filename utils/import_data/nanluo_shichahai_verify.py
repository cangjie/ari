#!/usr/bin/env python3
"""只读复核本轮工作簿、源文件版本、调用计量与存量业务行数。"""
from __future__ import annotations
import hashlib
import json
from collections import Counter
from decimal import Decimal
from pathlib import Path
from openpyxl import load_workbook
import meta_lib

HERE=Path(__file__).resolve().parent
KEY='nanluo_shichahai'
BASELINE={'beihai':23,'capital':6159,'ham':204,'jlpm':17723,'mfa_boston':203,'mfa_boston_ext':4464,'nmc':365,'palace':1757,'pem':467,'wmhg':61}
def load(suffix):return json.loads((HERE/(KEY+suffix)).read_text('utf-8'))
def digest(path):return hashlib.sha256(path.read_bytes()).hexdigest()
def dump(suffix,obj):
    (HERE/(KEY+suffix)).write_text(json.dumps(obj,ensure_ascii=False,indent=2,default=lambda x:int(x) if isinstance(x,Decimal) else str(x))+'\n',encoding='utf-8')

def main():
    data=load('_site_data.json');ratings=load('_ratings.json');report=load('_export_report.json')
    source_sha=digest(HERE/(KEY+'_site_data.json'))
    assert ratings['input_sha256']==load('_audit_manifest.json')['input_sha256']==load('_translations.json')['input_sha256']==source_sha
    assert ratings['model']=='gpt-5.6-luna effort=high/medium/xhigh (codex-cli)'
    assert load('_audit_manifest.json')['model']=='claude-sonnet-5 effort=medium (claude-cli)'
    audited={r['seq']:r for r in map(json.loads,(HERE/(KEY+'_audit1_slim.jsonl')).read_text().splitlines())}
    translations=load('_translations.json')['values']
    scored={r['seq']:r for r in ratings['results']};wanted={r['source_seq']:r for r in data['rows']}
    assert set(audited)==set(scored)==set(wanted)
    assert report['audit_state']=='complete' and report['audited_records']==len(wanted)
    assert load('_audit_manifest.json')['audit_packets_sha256']==digest(HERE/(KEY+'_audit_packets.json'))
    binding=load('_audit_resume_binding.json')
    assert load('_audit_manifest.json')['reused_audit_binding_sha256']==digest(HERE/(KEY+'_audit_resume_binding.json'))
    assert binding['current_rating_sha256']==digest(HERE/(KEY+'_ratings.json'))
    for rec in binding['reused_records']:
        seq=rec['seq']
        row_sha=lambda obj:hashlib.sha256(json.dumps(obj,ensure_ascii=False,sort_keys=True,separators=(',',':')).encode()).hexdigest()
        assert row_sha(scored[seq])==rec['rating_sha256']
        assert row_sha(audited[seq])==rec['audit_record_sha256']
    assert '郭沫若' in scored[48]['cr_reason'] and '郭守敬' not in scored[48]['cr_reason']
    path=HERE/'exports'/'展品_北京南锣鼓巷及什刹海.xlsx'
    wb=load_workbook(path,data_only=False)
    assert wb.sheetnames==['zh-CN','en','特色商铺','来源证据','同类组','说明']
    for name in ['zh-CN','en']:
        ws=wb[name]
        assert ws.max_row==len(wanted)+1
        seen=set()
        for row in ws.iter_rows(min_row=2,values_only=True):
            seq=row[0];assert seq in wanted and seq not in seen;seen.add(seq)
            assert row[5]==scored[seq]['tier']
            assert isinstance(row[6],(int,float)) and abs(row[6]-scored[seq]['core'])<=0.000051
            assert all(isinstance(x,(int,float)) and 0<=x<=10 for x in row[6:14])
            shown=wanted[seq].get('display_name',wanted[seq]['name'])
            if name=='zh-CN':assert row[1]==shown
            else:assert row[1]==translations[shown]['en']
            a=audited[seq];flag=a['found_factual_error'] or any(x['tier_sensitive'] for x in a['missing'])
            assert row[14]==(('是' if flag else '否') if name=='zh-CN' else ('Yes' if flag else 'No'))
            assert row[21]=='claude-sonnet-5 effort=medium (claude-cli)'
        assert seen==set(wanted)
        # 同时冻结表头和三列身份信息。
        assert ws.freeze_panes=='D2',(name,ws.freeze_panes)
        assert len(ws.tables)==1
        assert next(iter(ws.tables.values())).autoFilter is not None
    assert wb['来源证据'].max_row==report['evidence_records']+1
    assert wb['同类组'].max_row==len(wanted)+1
    shops={s:r for s,r in wanted.items() if r.get('entry_kind')=='特色商铺'}
    assert len(shops)==report['added_shops']==16
    assert wb['特色商铺'].max_row==len(shops)+1
    assert wb['特色商铺'].freeze_panes=='C2'
    for row in wb['特色商铺'].iter_rows(min_row=2,values_only=True):
        seq=row[0];assert seq in shops
        assert row[1]==shops[seq]['name'] and row[4]==scored[seq]['tier']
        assert abs(row[5]-scored[seq]['core'])<=0.000051
        assert row[7]==shops[seq]['specialty'] and row[8]==shops[seq]['address']
    assert wb['zh-CN']['B49'].value=='郭沫若故居（郭沫若纪念馆）'
    assert sum('郭沫若' in str(row[0]) for row in wb['zh-CN'].iter_rows(min_row=2,min_col=2,max_col=2,values_only=True))==1
    for ws in wb:
        for row in ws:
            for cell in row:
                assert cell.data_type!='e',(ws.title,cell.coordinate,cell.value)
    actual=Counter(row[0] for row in wb['zh-CN'].iter_rows(min_row=2,min_col=6,max_col=6,values_only=True))
    assert dict(actual)==report['tier']
    wb.close()
    conn=meta_lib.connect(autocommit=True)
    with conn.cursor() as q:
        q.execute('SELECT m.key_name,COUNT(a.id) FROM museum m LEFT JOIN artwork a ON a.museum_id=m.id GROUP BY m.key_name ORDER BY m.key_name')
        after=dict(q.fetchall());assert after==BASELINE,(after,BASELINE)
        q.execute('SELECT provider,model,effort,stage,status,COUNT(*),SUM(prompt_tokens),SUM(cached_tokens),SUM(completion_tokens),SUM(reasoning_tokens),SUM(latency_ms) FROM llm_call WHERE museum_key=%s GROUP BY provider,model,effort,stage,status ORDER BY provider,stage,effort',(KEY,))
        fields=['provider','model','effort','stage','status','calls','prompt_tokens','cached_tokens','completion_tokens','reasoning_tokens','latency_ms']
        usage=[dict(zip(fields,row)) for row in q.fetchall()]
        assert {r['model'] for r in usage}<={'gpt-5.6-luna','claude-sonnet-5'}
    summary={'records':len(wanted),'tier':dict(actual),'workbook':str(path),'workbook_bytes':path.stat().st_size,'workbook_sha256':digest(path),'source_sha256':source_sha,
             'bilingual_rows_equal':True,'typed_scores':True,'frozen_headers_and_identifiers':True,'business_before':BASELINE,'business_after':after,'business_total':sum(after.values()),
             'models':{'information':'gpt-5.6-luna medium','rating':ratings['model'],'audit':load('_audit_manifest.json')['model']},'audit_needs_review':report['needs_review'],'factual_errors':report['factual_errors'],
             'english_cjk_residual':report['english_cjk_residual'],'translation_revisions':len(load('_translation_repair.json')['changes']),
             'added_shops':len(shops),'shop_tier':report['shop_tier'],'guomoruo':report['guomoruo'],'changed_existing_tiers':report['changed_existing_tiers'],
             'reaudited_after_cr_correction':binding['reaudit_seqs'],'unchanged_audit_records_reused':len(binding['reused_records']),
             'limits':['来源与形式复核不等于现场验收；商铺营业、价格、菜单和体验均未现场逐店核实，平台/索引可能滞后。','Claude标准缓存的输入侧token有历史计量缺口；provider原始usage另留。','初版一次英文修订校验失败及一次Claude额度上限均保留；不能声称完整零失败计量。增补分组复核前主动结束的一轮，已返回的两组结果另存且保留模型缓存，未用于最终分组。']}
    if (HERE/(KEY+'_tour_audit_failed_session.json')).exists():
        summary['limits'].append('发布首轮CLI因大输入后的上下文压缩与结构字段丢失失败（llm_call.id=3928），原始会话和usage单独留存；标准缓存错误行的token仍为空，不补造成完整账单。')
    dump('_usage_report.json',dict(calls=usage,success_calls=sum(r['calls'] for r in usage if r['status']=='ok'),error_calls=sum(r['calls'] for r in usage if r['status']=='error'),limitations=summary['limits'][1:]))
    dump('_verification_report.json',summary)
    print(json.dumps(summary,ensure_ascii=False,indent=2))

if __name__=='__main__':main()
