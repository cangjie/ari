#!/usr/bin/env python3
# -*- coding: utf-8 -*-
"""
把 500cities_tier_v2.xlsx 导入 city / cultural_site / cultural_site_tier_change 三张表，
并把其中所有展示文本按内容ID拆进 content / content_text 两张多语种表。

用法：
  # 正式导入 MySQL
  python import_data.py --excel 500cities_tier_v2.xlsx \
      --host 44.207.251.65 --port 3306 --user ari --password *** --database ari

  # 不连 MySQL，用 SQLite 试跑一遍验证清洗逻辑
  python import_data.py --excel 500cities_tier_v2.xlsx --sqlite ./check.db

依赖：openpyxl；导入 MySQL 时另需 PyMySQL
"""

import argparse
import csv
import os
import re
import sys
from collections import Counter

import openpyxl

SHEET_CITY = "sheet1_城市列表"
SHEET_SITE = "sheet2_博物馆遗产地列表"
SHEET_CHANGE = "sheet3_等级变动明细"

TIER_ORDER = {"S": 4, "A": 3, "B": 2, "C": 1}

# ---------------------------------------------------------------- 多语种
LANG_ZH = "zh-CN"
LANG_EN = "en"

# content.kind 的取值与写入顺序。内容ID按这个顺序 + 文本排序确定性分配，
# 同一份 Excel 重跑得到同一批ID，便于 diff 与排查。
#
# 这批 kind 同时界定「本导入器拥有哪些内容」：content 表与 import_artworks.py
# 共用，清空时只能删自己这几种，否则会把展品文本一起删掉（而展品外键是
# RESTRICT，实际后果是本次导入直接失败）。
CONTENT_KINDS = (
    "city_name", "country_name", "site_name", "tier_reason",
    "change_reason", "collection_type", "data_source",
)

# 本数据集的内容ID段起点。展品数据集从 1,000,000 起，两段不重叠，
# 因此双方都能显式分配ID而不必担心撞号
CONTENT_ID_BASE = 0

# 译名表 confidence -> content_text.source
CONF_TO_SOURCE = {"官方": "人工校对", "AI": "AI翻译", "存疑": "存疑"}

RE_HAN = re.compile(r"[㐀-䶿一-鿿]")

# 「年参观人数/说明」列的判定：能解析成数字的是人次，否则是门类文本
RE_VISITS = re.compile(r"^([\d,]+)\s*(?:\((\d{4})\))?$")


# ---------------------------------------------------------------- 清洗辅助
def s(v):
    """转成去空白的字符串，空值返回 None。"""
    if v is None:
        return None
    v = str(v).replace(" ", " ").strip()
    return v or None


def parse_visitors(raw):
    """
    原表把两种语义塞进了同一列：
        '9,046,000 (2025)'      -> 年参观人次 + 年份
        '1,840,000'             -> 只有人次，无年份
        'Art & Encyclopedic'    -> 其实是门类
        None                    -> 无数据（577/825 条）
    返回 (annual_visitors, visitors_year, collection_type)
    """
    raw = s(raw)
    if raw is None:
        return None, None, None
    m = RE_VISITS.match(raw)
    if m:
        return int(m.group(1).replace(",", "")), (int(m.group(2)) if m.group(2) else None), None
    return None, None, raw


def parse_delta(raw):
    """'↑升 2 档' -> 2 ; '↓降 1 档' -> -1 ; '—' -> 0"""
    raw = s(raw)
    if not raw or raw in ("—", "-", "–"):
        return 0
    m = re.search(r"(\d+)", raw)
    n = int(m.group(1)) if m else 1
    return -n if ("降" in raw or "↓" in raw) else n


def delta_from_tiers(tier_from, tier_to):
    """用等级本身算档数，作为 Excel 文本的交叉校验。"""
    if not tier_from or not tier_to:
        return 0
    return TIER_ORDER[tier_to] - TIER_ORDER[tier_from]


def num(v, default=None):
    if v is None or str(v).strip() == "":
        return default
    return float(v)


def load_overrides(path):
    """
    读取双属性覆盖表 -> {(名称, 城市): (is_museum, is_heritage_site, confidence)}

    Excel 只有 track 这一个二选一字段，无法表达「既是博物馆又是遗址」。
    因此 is_museum / is_heritage_site 先按 track 打底，再用这张人工判定表覆盖。
    以 # 开头的行是注释。
    """
    if not path or not os.path.exists(path):
        return {}
    out = {}
    with open(path, encoding="utf-8") as f:
        rows = [ln for ln in f if not ln.lstrip().startswith("#")]
    for r in csv.DictReader(rows):
        if not r.get("name"):
            continue
        out[(r["name"].strip(), r["city"].strip())] = (
            int(r["is_museum"]), int(r["is_heritage_site"]), (r.get("confidence") or "").strip()
        )
    return out


def detect_lang(text):
    """判定一段原文的语种。含汉字即视为中文 —— 「文献展(documenta)」这类
    中英混排的条目属中文语境，其纯英文名由译名表另行给出。"""
    return LANG_ZH if RE_HAN.search(text) else LANG_EN


def load_translations(paths, kinds=CONTENT_KINDS):
    """
    读取译名表 -> {(kind, key): (zh, en, confidence)}

    列格式：kind,key,zh,en,confidence
      kind       —— content.kind，决定这段文本属于哪一类
      key        —— 原文，去重键，与主表 *_key 或原字段值逐字对齐
      zh / en    —— 两个语种的文本，其中一列等于 key（原文那一侧）
      confidence —— 官方 / AI / 存疑，只描述「译出来的那一半」
    以 # 开头的行是注释。

    kinds 是本调用方认得的 kind 集合。import_artworks.py 必须传自己那批，
    否则展品译名会被当成未知 kind 整批丢掉（而且只留一行 warn，很容易漏看）。
    """
    out = {}
    for path in paths:
        if not path or not os.path.exists(path):
            continue
        with open(path, encoding="utf-8") as f:
            rows = [ln for ln in f if not ln.lstrip().startswith("#")]
        for r in csv.DictReader(rows):
            kind, key = (r.get("kind") or "").strip(), (r.get("key") or "").strip()
            if not kind or not key:
                continue
            if kind not in kinds:
                print(f"  [warn] 译名表 {os.path.basename(path)} 有未知 kind：{kind}，已跳过")
                continue
            out[(kind, key)] = (
                (r.get("zh") or "").strip() or None,
                (r.get("en") or "").strip() or None,
                (r.get("confidence") or "").strip(),
            )
    return out


# ---------------------------------------------------------------- 读取 Excel
def read_excel(path, overrides=None):
    overrides = overrides or {}
    wb = openpyxl.load_workbook(path, read_only=True, data_only=True)

    # --- 城市 ---
    cities = []
    for r in wb[SHEET_CITY].iter_rows(min_row=2, values_only=True):
        if r[1] is None:
            continue
        cities.append(
            dict(
                city_rank=int(r[0]),
                name=s(r[1]),
                country=s(r[2]),
                continent=s(r[3]),
                score_art_museum=num(r[4], 0),
                score_history=num(r[5], 0),
                score_archaeology=num(r[6], 0),
                score_architecture=num(r[7], 0),
                score_accessibility=num(r[8], 0),
                score_total=num(r[9], 0),
                category=s(r[12]),
                flagship_name=s(r[11]),   # 仅用于回填，不直接入库
                site_count=int(r[10]),    # 仅用于校验，派生值不入库
            )
        )

    # --- 文化点位 ---
    sites = []
    for r in wb[SHEET_SITE].iter_rows(min_row=2, values_only=True):
        if r[5] is None:
            continue
        visitors, vyear, ctype = parse_visitors(r[17])
        tier, tier_prev = s(r[1]), s(r[2])
        delta_txt = parse_delta(r[3])
        delta_calc = delta_from_tiers(tier_prev, tier)
        if delta_txt != delta_calc:
            print(f"  [warn] 等级变动文本与等级不一致：{s(r[5])} 文本={delta_txt} 实算={delta_calc}，以实算为准")

        name, city_name, track = s(r[5]), s(r[6]), s(r[4])
        # 两个事实属性：先按 track 打底，再用人工判定表覆盖双属性条目
        is_museum = 1 if track == "博物馆" else 0
        is_heritage = 1 if track == "遗址地" else 0
        ov = overrides.get((name, city_name))
        if ov:
            is_museum, is_heritage, _ = ov

        sites.append(
            dict(
                name=name,
                city_name=city_name,
                country=s(r[7]),
                track=track,
                is_museum=is_museum,
                is_heritage_site=is_heritage,
                tier=tier,
                tier_prev=tier_prev,
                tier_delta=delta_calc,
                score_total=num(r[9]),
                dim1=num(r[10]), dim2=num(r[11]), dim3=num(r[12]),
                dim4=num(r[13]), dim5=num(r[14]), dim6=num(r[15]),
                tier_reason=s(r[16]),
                annual_visitors=visitors,
                visitors_year=vyear,
                collection_type=ctype,
                data_source=s(r[18]),
            )
        )

    # --- 等级变动明细 ---
    changes = []
    for r in wb[SHEET_CHANGE].iter_rows(min_row=2, values_only=True):
        if r[4] is None:
            continue
        tier_from, tier_to = s(r[0]), s(r[1])
        changes.append(
            dict(
                name=s(r[4]),
                city_name=s(r[5]),
                tier_from=tier_from,
                tier_to=tier_to,
                tier_delta=delta_from_tiers(tier_from, tier_to),
                score_total=num(r[7]),
                change_reason=s(r[8]),
            )
        )

    wb.close()
    return cities, sites, changes


# ---------------------------------------------------------------- 数据自检
def sanity_check(cities, sites, changes):
    problems = []

    ranks = [c["city_rank"] for c in cities]
    if len(set(ranks)) != len(ranks):
        problems.append("city_rank 有重复")
    keys = [(c["name"], c["country"]) for c in cities]
    if len(set(keys)) != len(keys):
        problems.append("(城市, 国家) 有重复")

    city_index = {c["name"] for c in cities}
    orphan = {x["city_name"] for x in sites if x["city_name"] not in city_index}
    if orphan:
        problems.append(f"{len(orphan)} 个点位所在城市不在城市表：{sorted(orphan)[:5]}")

    sk = [(x["name"], x["city_name"]) for x in sites]
    dup = [k for k, n in Counter(sk).items() if n > 1]
    if dup:
        problems.append(f"(名称, 城市) 有重复：{dup[:5]}")

    skset = set(sk)
    miss = [(c["name"], c["city_name"]) for c in changes if (c["name"], c["city_name"]) not in skset]
    if miss:
        problems.append(f"{len(miss)} 条变动明细找不到对应点位：{miss[:5]}")

    noattr = [x["name"] for x in sites if not x["is_museum"] and not x["is_heritage_site"]]
    if noattr:
        problems.append(f"{len(noattr)} 条 is_museum 与 is_heritage_site 同时为 0：{noattr[:5]}")

    # 城市表的「博物馆遗产地数量」与实际条数对不上——原表已知问题，只提示不阻断
    actual = Counter(x["city_name"] for x in sites)
    mismatch = [(c["name"], c["site_count"], actual.get(c["name"], 0))
                for c in cities if c["site_count"] != actual.get(c["name"], 0)]
    if mismatch:
        print(f"  [info] {len(mismatch)} 个城市的「博物馆遗产地数量」与实际行数不符（原表逗号拆行遗留）：{mismatch[:5]}")

    return problems


# ---------------------------------------------------------------- 写库
class Target:
    """屏蔽 MySQL / SQLite 的占位符差异。"""

    def __init__(self, conn, flavor):
        self.conn, self.flavor = conn, flavor
        self.ph = "%s" if flavor == "mysql" else "?"

    def exec(self, sql, args=None):
        cur = self.conn.cursor()
        cur.execute(sql, args or ())
        return cur

    def many(self, sql, rows):
        cur = self.conn.cursor()
        cur.executemany(sql, rows)
        return cur

    def insert(self, table, cols):
        return f"INSERT INTO {table} ({', '.join(cols)}) VALUES ({', '.join([self.ph] * len(cols))})"


CITY_COLS = ["name_key", "name_cid", "country_key", "country_cid", "continent",
             "score_art_museum", "score_history",
             "score_archaeology", "score_architecture", "score_accessibility",
             "score_total", "city_rank", "category"]

SITE_COLS = ["city_id", "name_key", "name_cid", "is_museum", "is_heritage_site", "track",
             "tier", "tier_prev", "tier_delta", "score_total",
             "dim1", "dim2", "dim3", "dim4", "dim5", "dim6", "tier_reason_cid",
             "annual_visitors", "visitors_year", "collection_type_cid", "data_source_cid"]

CHANGE_COLS = ["site_id", "rule_version_from", "rule_version_to", "tier_from", "tier_to",
               "tier_delta", "score_total", "change_reason_cid"]


# ---------------------------------------------------------------- 内容表
def collect_contents(cities, sites, changes):
    """把三个数据集里所有可翻译文本按 kind 去重收集 -> {kind: [原文, ...]}"""
    buckets = {k: set() for k in CONTENT_KINDS}
    for c in cities:
        buckets["city_name"].add(c["name"])
        buckets["country_name"].add(c["country"])
    for x in sites:
        buckets["site_name"].add(x["name"])
        for kind, field in (("tier_reason", "tier_reason"),
                            ("collection_type", "collection_type"),
                            ("data_source", "data_source")):
            if x[field]:
                buckets[kind].add(x[field])
    for ch in changes:
        if ch["change_reason"]:
            buckets["change_reason"].add(ch["change_reason"])
    # 排序而非按出现顺序：同一份 Excel 即使行序变了，内容ID也不变
    return {k: sorted(v) for k, v in buckets.items()}


def build_content_rows(buckets, trans, id_base=CONTENT_ID_BASE):
    """
    为每段去重原文分配内容ID，并生成 content / content_text 两组待插入行。

    内容ID显式指定而非依赖 AUTO_INCREMENT：2416 段若逐条插入再取 lastrowid，
    走公网就是 2416 次往返；显式分配可以 executemany 一次灌完，
    而且同一份数据重跑得到同一批ID。

    id_base 是本数据集的ID段起点，用来与另一个导入器错开号段。

    返回 (cid_of, content_rows, text_rows, stats)
    """
    cid_of, content_rows, text_rows = {}, [], []
    stats = Counter()
    next_id = id_base
    for kind in CONTENT_KINDS:
        for key in buckets[kind]:
            next_id += 1
            cid_of[(kind, key)] = next_id
            content_rows.append((next_id, kind))

            zh, en, conf = trans.get((kind, key), (None, None, ""))
            src = CONF_TO_SOURCE.get(conf, "AI翻译")
            # 原文永远保底写入，即使译名表缺这一条
            texts = {detect_lang(key): (key, "原始")}
            if zh:
                texts[LANG_ZH] = (zh, "原始" if zh == key else src)
            if en:
                texts[LANG_EN] = (en, "原始" if en == key else src)
            for lang, (text, source) in texts.items():
                text_rows.append((next_id, lang, text, source))
                stats[source] += 1
            if len(texts) < 2:
                stats["缺译"] += 1
    return cid_of, content_rows, text_rows, stats


def load(t, cities, sites, changes, trans):
    # 清空（按外键顺序：先解开 city -> site 的回填指针，再从叶子往根删，内容表最后）
    if t.flavor == "mysql":
        t.exec("UPDATE city SET flagship_site_id = NULL")
    for tb in ("cultural_site_tier_change", "cultural_site", "city"):
        t.exec(f"DELETE FROM {tb}")
    # 内容表只删本导入器名下的 kind。整表清空会连 import_artworks.py 写的展品文本
    # 一起删掉 —— 展品外键是 RESTRICT，真删起来这一句就会直接报错
    kinds = ", ".join(f"'{k}'" for k in CONTENT_KINDS)
    t.exec(f"DELETE FROM content_text WHERE content_id IN "
           f"(SELECT id FROM content WHERE kind IN ({kinds}))")
    t.exec(f"DELETE FROM content WHERE kind IN ({kinds})")
    if t.flavor == "mysql":
        # content 不重置 AUTO_INCREMENT：ID 已按号段显式分配，重置只会干扰另一个数据集
        for tb in ("cultural_site_tier_change", "cultural_site", "city"):
            t.exec(f"ALTER TABLE {tb} AUTO_INCREMENT = 1")

    # 内容表：所有展示文本先落地，主表随后只引用内容ID
    buckets = collect_contents(cities, sites, changes)
    content_id_of, content_rows, text_rows, stats = build_content_rows(buckets, trans)
    t.many(t.insert("content", ["id", "kind"]), content_rows)
    t.many(t.insert("content_text", ["content_id", "lang", "text", "source"]), text_rows)
    print(f"  content                   {len(content_rows):>4} 段"
          f"（{'、'.join(f'{k} {len(v)}' for k, v in buckets.items() if v)}）")
    print(f"  content_text              {len(text_rows):>4} 行"
          f"（原始 {stats['原始']} / 人工校对 {stats['人工校对']}"
          f" / AI翻译 {stats['AI翻译']} / 存疑 {stats['存疑']}）")
    if stats["缺译"]:
        print(f"  [info] {stats['缺译']} 段只有一个语种，展示时需按 locale 回落")

    # 城市
    for c in cities:
        c["name_key"], c["name_cid"] = c["name"], content_id_of[("city_name", c["name"])]
        c["country_key"] = c["country"]
        c["country_cid"] = content_id_of[("country_name", c["country"])]
    t.many(t.insert("city", CITY_COLS), [tuple(c[k] for k in CITY_COLS) for c in cities])
    city_id = {nk: i for i, nk in t.exec("SELECT id, name_key FROM city").fetchall()}
    print(f"  city                      {len(city_id):>4} 行")

    # 文化点位
    rows = []
    for x in sites:
        x["city_id"] = city_id[x["city_name"]]
        x["name_key"], x["name_cid"] = x["name"], content_id_of[("site_name", x["name"])]
        for kind, field in (("tier_reason", "tier_reason"),
                            ("collection_type", "collection_type"),
                            ("data_source", "data_source")):
            x[field + "_cid"] = content_id_of[(kind, x[field])] if x[field] else None
        rows.append(tuple(x[k] for k in SITE_COLS))
    t.many(t.insert("cultural_site", SITE_COLS), rows)
    site_id = {(nk, ci): sid
               for sid, nk, ci in
               t.exec("SELECT id, name_key, city_id FROM cultural_site").fetchall()}
    print(f"  cultural_site             {len(site_id):>4} 行")

    # 变动明细
    rows = []
    for c in changes:
        sid = site_id[(c["name"], city_id[c["city_name"]])]
        reason_cid = content_id_of[("change_reason", c["change_reason"])] if c["change_reason"] else None
        rows.append((sid, "v1", "v2", c["tier_from"], c["tier_to"],
                     c["tier_delta"], c["score_total"], reason_cid))
    t.many(t.insert("cultural_site_tier_change", CHANGE_COLS), rows)
    print(f"  cultural_site_tier_change {len(rows):>4} 行")

    # 回填城市的代表性点位
    ph = t.ph
    filled = unmatched = 0
    for c in cities:
        ci = city_id[c["name"]]
        sid = site_id.get((c["flagship_name"], ci))
        if sid is None and c["flagship_name"]:
            # 原表有 7 处名称因含中文逗号被截断，退化为前缀匹配
            row = t.exec(
                f"SELECT id FROM cultural_site WHERE city_id = {ph} AND name_key LIKE {ph} "
                f"ORDER BY score_total DESC",
                (ci, c["flagship_name"] + "%"),
            ).fetchone()
            sid = row[0] if row else None
        if sid is None:
            unmatched += 1
            continue
        t.exec(f"UPDATE city SET flagship_site_id = {ph} WHERE id = {ph}", (sid, ci))
        filled += 1
    print(f"  flagship_site_id 回填 {filled} 个城市" + (f"，{unmatched} 个未匹配" if unmatched else ""))

    t.conn.commit()


# ---------------------------------------------------------------- SQLite 试跑
SQLITE_DDL = """
PRAGMA foreign_keys = ON;
CREATE TABLE IF NOT EXISTS content (
  id INTEGER PRIMARY KEY,
  kind TEXT NOT NULL CHECK (kind IN ('city_name','country_name','site_name','tier_reason',
                                     'change_reason','collection_type','data_source'))
);
CREATE TABLE IF NOT EXISTS content_text (
  content_id INTEGER NOT NULL REFERENCES content(id) ON DELETE CASCADE,
  lang TEXT NOT NULL,
  text TEXT NOT NULL,
  source TEXT NOT NULL DEFAULT '原始'
         CHECK (source IN ('原始','AI翻译','存疑','人工校对')),
  PRIMARY KEY (content_id, lang)
);
CREATE TABLE IF NOT EXISTS city (
  id INTEGER PRIMARY KEY AUTOINCREMENT,
  name_key TEXT NOT NULL,
  name_cid INTEGER NOT NULL REFERENCES content(id),
  country_key TEXT NOT NULL,
  country_cid INTEGER NOT NULL REFERENCES content(id),
  continent TEXT NOT NULL CHECK (continent IN ('亚洲','欧洲','非洲','北美洲','南美洲','大洋洲')),
  score_art_museum REAL NOT NULL DEFAULT 0, score_history REAL NOT NULL DEFAULT 0,
  score_archaeology REAL NOT NULL DEFAULT 0, score_architecture REAL NOT NULL DEFAULT 0,
  score_accessibility REAL NOT NULL DEFAULT 0, score_total REAL NOT NULL DEFAULT 0,
  city_rank INTEGER NOT NULL UNIQUE,
  category TEXT NOT NULL CHECK (category IN ('综合类','艺术博物馆','考古遗址','历史建筑群','宗教遗产')),
  flagship_site_id INTEGER REFERENCES cultural_site(id),
  UNIQUE (name_key, country_key),
  CHECK (score_total BETWEEN 0 AND 100)
);
CREATE TABLE IF NOT EXISTS cultural_site (
  id INTEGER PRIMARY KEY AUTOINCREMENT,
  city_id INTEGER NOT NULL REFERENCES city(id),
  name_key TEXT NOT NULL,
  name_cid INTEGER NOT NULL REFERENCES content(id),
  is_museum INTEGER NOT NULL DEFAULT 0 CHECK (is_museum IN (0,1)),
  is_heritage_site INTEGER NOT NULL DEFAULT 0 CHECK (is_heritage_site IN (0,1)),
  track TEXT NOT NULL CHECK (track IN ('博物馆','遗址地')),
  tier TEXT NOT NULL CHECK (tier IN ('S','A','B','C')),
  tier_prev TEXT CHECK (tier_prev IN ('S','A','B','C')),
  tier_delta INTEGER NOT NULL DEFAULT 0 CHECK (tier_delta BETWEEN -3 AND 3),
  score_total REAL NOT NULL CHECK (score_total BETWEEN 0 AND 100),
  dim1 REAL NOT NULL CHECK (dim1 BETWEEN 0 AND 10), dim2 REAL NOT NULL CHECK (dim2 BETWEEN 0 AND 10),
  dim3 REAL NOT NULL CHECK (dim3 BETWEEN 0 AND 10), dim4 REAL NOT NULL CHECK (dim4 BETWEEN 0 AND 10),
  dim5 REAL NOT NULL CHECK (dim5 BETWEEN 0 AND 10), dim6 REAL NOT NULL CHECK (dim6 BETWEEN 0 AND 10),
  tier_reason_cid INTEGER REFERENCES content(id),
  annual_visitors INTEGER, visitors_year INTEGER,
  collection_type_cid INTEGER REFERENCES content(id),
  data_source_cid INTEGER REFERENCES content(id),
  UNIQUE (name_key, city_id),
  CHECK (is_museum = 1 OR is_heritage_site = 1)
);
CREATE TABLE IF NOT EXISTS cultural_site_tier_change (
  id INTEGER PRIMARY KEY AUTOINCREMENT,
  site_id INTEGER NOT NULL REFERENCES cultural_site(id) ON DELETE CASCADE,
  rule_version_from TEXT NOT NULL DEFAULT 'v1', rule_version_to TEXT NOT NULL DEFAULT 'v2',
  tier_from TEXT NOT NULL, tier_to TEXT NOT NULL,
  tier_delta INTEGER NOT NULL CHECK (tier_delta BETWEEN -3 AND 3 AND tier_delta <> 0),
  score_total REAL,
  change_reason_cid INTEGER REFERENCES content(id),
  UNIQUE (site_id, rule_version_from, rule_version_to),
  CHECK (tier_from <> tier_to)
);
CREATE VIEW IF NOT EXISTS v_cultural_site_full AS
SELECT s.id, c.city_rank, s.tier, s.tier_prev, s.tier_delta, s.track,
       s.is_museum, s.is_heritage_site,
       s.name_key, sn_zh.text AS name_zh, sn_en.text AS name_en,
       c.name_key AS city_key, cn_zh.text AS city_zh, cn_en.text AS city_en,
       c.country_key, co_zh.text AS country_zh, co_en.text AS country_en,
       c.continent, s.score_total,
       s.dim1, s.dim2, s.dim3, s.dim4, s.dim5, s.dim6,
       tr_zh.text AS tier_reason_zh, tr_en.text AS tier_reason_en,
       s.annual_visitors, s.visitors_year,
       cl_zh.text AS collection_type_zh, cl_en.text AS collection_type_en,
       ds_zh.text AS data_source_zh, ds_en.text AS data_source_en
FROM cultural_site s
JOIN city c ON c.id = s.city_id
LEFT JOIN content_text sn_zh ON sn_zh.content_id = s.name_cid            AND sn_zh.lang = 'zh-CN'
LEFT JOIN content_text sn_en ON sn_en.content_id = s.name_cid            AND sn_en.lang = 'en'
LEFT JOIN content_text cn_zh ON cn_zh.content_id = c.name_cid            AND cn_zh.lang = 'zh-CN'
LEFT JOIN content_text cn_en ON cn_en.content_id = c.name_cid            AND cn_en.lang = 'en'
LEFT JOIN content_text co_zh ON co_zh.content_id = c.country_cid         AND co_zh.lang = 'zh-CN'
LEFT JOIN content_text co_en ON co_en.content_id = c.country_cid         AND co_en.lang = 'en'
LEFT JOIN content_text tr_zh ON tr_zh.content_id = s.tier_reason_cid     AND tr_zh.lang = 'zh-CN'
LEFT JOIN content_text tr_en ON tr_en.content_id = s.tier_reason_cid     AND tr_en.lang = 'en'
LEFT JOIN content_text cl_zh ON cl_zh.content_id = s.collection_type_cid AND cl_zh.lang = 'zh-CN'
LEFT JOIN content_text cl_en ON cl_en.content_id = s.collection_type_cid AND cl_en.lang = 'en'
LEFT JOIN content_text ds_zh ON ds_zh.content_id = s.data_source_cid     AND ds_zh.lang = 'zh-CN'
LEFT JOIN content_text ds_en ON ds_en.content_id = s.data_source_cid     AND ds_en.lang = 'en';
"""


def main():
    ap = argparse.ArgumentParser()
    ap.add_argument("--excel", required=True)
    ap.add_argument("--overrides", default="dual_attribute_overrides.csv",
                    help="双属性人工判定表；不存在则两个布尔字段全部按 track 打底")
    ap.add_argument("--translations", nargs="*",
                    default=["translations_city.csv", "translations_site.csv",
                             "translations_text.csv"],
                    help="译名表；缺失的语种展示时靠回落，不阻断导入")
    ap.add_argument("--sqlite", help="试跑模式：写入指定 SQLite 文件而非 MySQL")
    ap.add_argument("--host", default="127.0.0.1")
    ap.add_argument("--port", type=int, default=3306)
    ap.add_argument("--user", default="ari")
    ap.add_argument("--password", default=os.environ.get("MYSQL_PASSWORD", ""))
    ap.add_argument("--database", default="ari")
    args = ap.parse_args()

    overrides = load_overrides(args.overrides)
    if overrides:
        n_doubt = sum(1 for v in overrides.values() if v[2] == "存疑")
        print(f"读取双属性判定表 … {len(overrides)} 条" + (f"（其中 {n_doubt} 条标记为存疑）" if n_doubt else ""))
    else:
        print(f"未找到 {args.overrides}，两个布尔字段按 track 打底")

    print("读取 Excel …")
    cities, sites, changes = read_excel(args.excel, overrides)
    print(f"  城市 {len(cities)} / 点位 {len(sites)} / 变动 {len(changes)}")

    # 覆盖表里若有名称/城市对不上 Excel 的行，静默丢弃会掩盖笔误，这里显式报出来
    site_keys = {(x["name"], x["city_name"]) for x in sites}
    stale = [k for k in overrides if k not in site_keys]
    if stale:
        print(f"  [warn] 判定表有 {len(stale)} 条未匹配到点位（名称或城市写错？）：{stale[:5]}")

    both = sum(1 for x in sites if x["is_museum"] and x["is_heritage_site"])
    print(f"  属性分布：仅博物馆 {sum(1 for x in sites if x['is_museum'] and not x['is_heritage_site'])}"
          f" / 仅遗址地 {sum(1 for x in sites if x['is_heritage_site'] and not x['is_museum'])}"
          f" / 双属性 {both}")

    # 译名表
    trans = load_translations(args.translations)
    buckets = collect_contents(cities, sites, changes)
    total_seg = sum(len(v) for v in buckets.values())
    if trans:
        n_doubt = sum(1 for v in trans.values() if v[2] == "存疑")
        print(f"读取译名表 … {len(trans)} 条 / 共 {total_seg} 段可翻译文本"
              + (f"（其中 {n_doubt} 条标记为存疑）" if n_doubt else ""))
    else:
        print(f"未找到译名表，{total_seg} 段文本只会有源语言一种")

    # 译名表若有对不上原文的行，静默丢弃会掩盖笔误，与判定表同样显式报出来
    known = {(kind, key) for kind, keys in buckets.items() for key in keys}
    stale_t = [k for k in trans if k not in known]
    if stale_t:
        print(f"  [warn] 译名表有 {len(stale_t)} 条未匹配到原文（kind 或 key 写错？）：{stale_t[:3]}")
    empty_t = [k for k, v in trans.items() if not v[0] and not v[1]]
    if empty_t:
        print(f"  [warn] 译名表有 {len(empty_t)} 条 zh 与 en 同时为空，等于没填：{empty_t[:3]}")

    print("数据自检 …")
    problems = sanity_check(cities, sites, changes)
    if problems:
        print("  发现问题，已中止：")
        for p in problems:
            print("   -", p)
        sys.exit(1)
    print("  通过")

    print("写入数据库 …")
    if args.sqlite:
        import sqlite3
        if os.path.exists(args.sqlite):
            os.remove(args.sqlite)
        conn = sqlite3.connect(args.sqlite)
        conn.executescript(SQLITE_DDL)
        load(Target(conn, "sqlite"), cities, sites, changes, trans)
        conn.close()
    else:
        import pymysql
        conn = pymysql.connect(host=args.host, port=args.port, user=args.user,
                               password=args.password, database=args.database,
                               charset="utf8mb4", autocommit=False)
        load(Target(conn, "mysql"), cities, sites, changes, trans)
        conn.close()
    print("完成")


if __name__ == "__main__":
    main()
