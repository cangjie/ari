#!/usr/bin/env python3
# -*- coding: utf-8 -*-
"""
从官网抓取数据确定性地生成北海公园的源 Excel，零网络、零 API。

    python3 beihai_build.py            # 写 artworks/北海公园_景点清单.xlsx + beihai_seq_map.csv + 体检报告
    python3 beihai_build.py --check    # 只做配对与体检，不写文件

输入（已提交）：beihai_site_data.json（beihai_site_scrape.py --scrape，2026-10-08 抓取）

**范围是用户定的**（2026-10-08），不在这里自行扩大：
· 中文站「景点介绍」（explore-jdjs）21 条全收；
· 英文站「景点介绍」（yingwen-jdjs）8 条里，中文站没有的 2 条（Jade Islet 琼华岛、White Dagoba 白塔）也收，
  共 23 行。其余 6 条与中文站逐条对上，作为该行的馆方英文（EN_PAIR）。
· 景点就是「展品」：一行一个景点，不另加北海公园整体一行。

**英文**（用户定）：英文站有的 8 条，英文名与英文简介用馆方英文，入库记「原始」；其余 15 条英文留空，
交给 translate_artwork.py 补译。**中文站标题里附带的英文不用** —— 那是机器直译（团城「Mission City」、
九龙壁「Kowloon wall」、阅古楼「unknow」），只留在抓取原文里。
英文站独有的 2 条中文名也留空，交给 translate_artwork.py --direction en2zh（同伪满皇宫九谷盘）。

**开放状态**（用户定）：只认官网原文写明「向游人开放」的（画舫斋、快雪堂、小西天），其余记未知。
官网列出 ≠ 开放（MFA「来源是官网 ≠ 在展」那条教训）。

**展厅列放「所在景点」**：V3.0 要求每个对象有父级范围（Tier Context）。只在官网原文写明某景点
在团城上、琼华岛上时才填（ZONE，逐条附原文并在构建时核对原文确实出现），写不出原文的留空 ——
阅古楼、铜仙承露盘实际都在琼华岛上，但官网这两条没写，就不填。

**序号按身份键幂等**（AGENTS.md 第 14 条）：身份键是官网文章 id（zh:<id> / en:<id>），
存在 beihai_seq_map.csv，已有的键永不改号，新键追加在末尾，消失的键留空号。
"""

from __future__ import annotations

import argparse
import csv
import datetime as dt
import html
import json
import pathlib
import re
import statistics
import sys
from collections import Counter

HERE = pathlib.Path(__file__).resolve().parent
BASE = "https://www.beihaipark.com.cn"
STATIC = BASE + "/api/sys/common/static/"        # 前端 window._CONFIG.staticDomainURL
DATA = HERE / "beihai_site_data.json"
XLSX = HERE / "artworks" / "北海公园_景点清单.xlsx"
SEQ_MAP = HERE / "beihai_seq_map.csv"
REPORT = HERE / "beihai_samples" / "build_report.txt"

ZH_COL = "explore-jdjs"
EN_COL = "yingwen-jdjs"

# 英文站 id -> (中文站 id, 凭什么是同一个景点：英文正文里的原话)
EN_PAIR = {
    "ab3ed1e3276d47f2a32012c1f0a414c4": ("a5a1a9ebfd9c4a1a8c4f0c28068e70a3", "The Round City (Tuancheng)"),
    "66ad74c150d5459198fa4fc66d8ca6a3": ("e61f3858b4654c42948bee513d1f733f", "The Temple of Eternal Peace (Yong’an Temple)"),
    "4ad695a57fb54e1c95e7cc07c5614961": ("3d515bb55280465c9bf0b6bb0f6bbb5b", "the Western Heaven Temple (also named Great Western Heaven)"),
    "7be5fbb37fac4fa8b0c1baad27ff4303": ("d9d392e933884262bff657fca64602a6", "inspired by the essay Painted Boat Hall by Song Dynasty writer Ouyang Xiu"),
    "3714fd4d4fcb4e5390eb071dd6b80c2b": ("26b8aa802284431c82905aed52581a75", "The Minor Western Heaven (Xiaoxitian)"),
    "eacdf5911e8542a5982f210e37b7f609": ("57ec180c97c844fba0ebbd17de5cacdf", "“Qiong Dao Chun Yin”（Jade Islet in Spring Shade）"),
}
# 英文站独有、中文站没有的
EN_ONLY = {
    "a99c01c264474b00a041a438926959c7": "Jade Islet（琼华岛）",
    "1d7c20bd4abc4595aa285bb303b1d97a": "White Dagoba（白塔）",
}

# 所在景点：中文站 id 或 en:<英文站 id> -> (所在景点, 官网原文, 原文出自哪一条)
ZONE = {
    "e7e394a842ae498ab7884c2aaa08abde": ("团城", "承光殿是团城的主要建筑", "本条"),                    # 承光殿
    "a1a20493b83b4ff28793958fdd8aa39f": ("团城", "院中有玉瓮亭，亭中有元代遗物玉瓮", "团城条"),        # 玉瓮
    "9C8802BA1A4249818ED0A95DDA5653FD": ("团城", "一年盛夏清乾隆皇帝来游团城，宫人摆案于树下", "本条"),  # 遮荫侯
    "0a5cc638e98f4f93834b69aaf9160234": ("团城", "守护在团城上", "本条"),                              # 白袍将军
    "e19c2fef4cca4b41895e7c89d17c3cdf": ("琼华岛", "坐落于北海公园琼华岛北侧", "本条"),               # 漪澜堂
    "57ec180c97c844fba0ebbd17de5cacdf": ("琼华岛", "立于琼华岛东坡", "本条"),                          # 琼岛春阴碑
    "8d8a18c6964743309fd6688a83af3a04": ("琼华岛", "位于琼岛东麓山脚下", "本条"),                      # 智珠殿
    "e61f3858b4654c42948bee513d1f733f": ("琼华岛", "Temple of Eternal Peace (Yong’an), a main building on the same "
                                                   "north-south axis as the dagoba, stands to the south of the White "
                                                   "Dagoba Hill", "英文站 Jade Islet 条"),         # 永安寺
    "en:1d7c20bd4abc4595aa285bb303b1d97a": ("琼华岛", "The Islet’s principal site is the Tibetan style White Dagoba",
                                            "英文站 Jade Islet 条"),                                # 白塔
}
# 原文出自哪一条 -> 那一条的 id（核对原文用）
ZONE_SRC = {"团城条": "a5a1a9ebfd9c4a1a8c4f0c28068e70a3", "英文站 Jade Islet 条": "a99c01c264474b00a041a438926959c7"}

# 开放状态：只认原文写明「向游人开放」的
ON_YES_PAT = re.compile(r"[^，。；]*(?:面向|向)游人开放[^，。；]*")
ON_NO = "开放状态未知（官网未写明）"


def _txt(s) -> str:
    """HTML -> 纯文本：段落间换行，去标签与图片，解实体，去每段开头的缩进空白。"""
    s = re.sub(r"<br\s*/?>|</p>|</div>", "\n", str(s or ""), flags=re.I)
    s = re.sub(r"<[^>]+>", "", s)
    s = html.unescape(s).replace("\xa0", " ").replace("　", " ")
    lines = [re.sub(r"[ \t]+", " ", ln).strip() for ln in s.split("\n")]
    return "\n".join(ln for ln in lines if ln)


def _flat(s: str) -> str:
    return re.sub(r"\s+", "", s or "")


def _split_zh(title: str) -> tuple[str, str]:
    """中文站标题「漪澜堂  Yi Lan Tang」-> (中文名, 那串机翻英文)。与前端的拆法相同。"""
    m = re.match(r"^([一-鿿]+)\s*(.*)$", title.strip())
    if not m:
        sys.exit(f"[fatal] 中文站标题不是「汉字 + 英文」的写法：{title!r}")
    return m.group(1), m.group(2).strip()


def _dedup_title(t: str) -> tuple[str, str]:
    """英文站有一条标题把自己重复了一遍（「Stele … Shade Stele … Shade」），去重并返回说明。"""
    t = re.sub(r"\s+", " ", t.strip())
    w = t.split(" ")
    if len(w) % 2 == 0 and w[:len(w) // 2] == w[len(w) // 2:]:
        half = " ".join(w[:len(w) // 2])
        return half, f"官网英文标题把自己重复了一遍（「{t}」），去重为「{half}」"
    return t, ""


def _desc(a: dict) -> str:
    """简介 + 正文：简介若已整句出现在正文里就只取正文，否则简介在前、正文在后。"""
    syn, body = _txt(a.get("fdSynopsis")), _txt(a.get("fdContent"))
    if not body:
        return syn
    if not syn or _flat(syn) in _flat(body):
        return body
    return syn + "\n" + body


def _img(a: dict) -> str:
    p = (a.get("fdPicture") or "").split(",")[0].strip()
    return STATIC + p if p else ""


def load() -> dict:
    d = json.loads(DATA.read_text("utf-8"))
    lists = {v["code"]: v for v in d["lists"].values() if v.get("code")}
    for code in (ZH_COL, EN_COL):
        if code not in lists or lists[code].get("incomplete") or lists[code].get("truncated"):
            sys.exit(f"[fatal] 栏目 {code} 没取全或不在抓取数据里")
    return d


def build_rows(d: dict, r) -> list[dict]:
    lists = {v["code"]: v for v in d["lists"].values() if v.get("code")}
    det = d["details"]
    zh_ids = [str(x["id"]) for x in lists[ZH_COL]["records"]]
    en_ids = [str(x["id"]) for x in lists[EN_COL]["records"]]
    if sorted(en_ids) != sorted(list(EN_PAIR) + list(EN_ONLY)):
        sys.exit(f"[fatal] 英文站景点与 EN_PAIR + EN_ONLY 登记的对不上：{sorted(en_ids)}")
    en_of = {}
    for eid, (zid, quote) in EN_PAIR.items():
        if zid not in zh_ids:
            sys.exit(f"[fatal] EN_PAIR 登记的中文站 id {zid} 不在景点介绍里")
        if quote not in _txt(det[eid].get("fdContent")):
            sys.exit(f"[fatal] EN_PAIR {eid} 的凭据原文不在英文正文里：{quote!r}")
        en_of[zid] = eid
    for k, (zone, quote, where) in ZONE.items():
        if where == "本条":
            src_id = k.split(":", 1)[1] if k.startswith("en:") else k
        else:
            src_id = ZONE_SRC[where]
        if _flat(quote) not in _flat(_desc(det[src_id]) + _txt(det[src_id].get("fdContent"))):
            sys.exit(f"[fatal] ZONE {k} 的原文没在 {where} 里找到：{quote!r}")

    rows: list[dict] = []
    for zid in zh_ids:
        a = det[zid]
        if a.get("fdIsExternalLinks") == "2":
            sys.exit(f"[fatal] 景点 {zid} 是外链，没有正文：{a.get('fdExternalLinksUrl')}")
        name, mt_en = _split_zh(a["fdTitle"])
        desc = _desc(a)
        eid = en_of.get(zid)
        name_en = desc_en = en_url = note_en = ""
        if eid:
            e = det[eid]
            name_en, note_en = _dedup_title(e["fdTitle"])
            desc_en = _txt(e.get("fdContent"))
            en_url = f"{BASE}/#/english/jdxq?id={eid}"
        on = ON_YES_PAT.search(desc)
        zone = ZONE.get(zid)
        rows.append(dict(
            key=f"zh:{zid}", hall=zone[0] if zone else "", name=name, desc=desc, image=_img(a),
            on=f"向游人开放（官网原文：「{on.group(0).strip()}」）" if on else ON_NO,
            url=f"{BASE}/#/detail/{zid}", level="景点", cat="官网景点介绍",
            name_en=name_en, desc_en=desc_en,
            oid=f"detail/{zid}" + (f" · en/{eid}" if eid else ""),
            src="beihai_official（中文站景点介绍）" + (" + 英文站" if eid else ""),
            en_url=en_url,
            note="；".join(x for x in (
                f"中文站标题附带的机翻英文「{mt_en}」未采用" if mt_en else "",
                f"所在景点据{zone[2]}原文：「{zone[1]}」" if zone else "",
                f"英文对应凭据：「{EN_PAIR[eid][1]}」" if eid else "", note_en) if x)))
    for eid, why in EN_ONLY.items():
        e = det[eid]
        name_en, note_en = _dedup_title(e["fdTitle"])
        zone = ZONE.get(f"en:{eid}")
        desc_en = _txt(e.get("fdContent"))
        on = ON_YES_PAT.search(desc_en)
        rows.append(dict(
            key=f"en:{eid}", hall=zone[0] if zone else "", name="", desc="", image=_img(e),
            on=ON_NO, url=f"{BASE}/#/english/jdxq?id={eid}", level="景点", cat="官网景点介绍（仅英文站）",
            name_en=name_en, desc_en=desc_en, oid=f"en/{eid}", src="beihai_official（仅英文站景点介绍）",
            en_url=f"{BASE}/#/english/jdxq?id={eid}",
            note="；".join(x for x in (
                f"中文站景点介绍里没有这一条：{why}；中文名与中文简介由 translate_artwork.py 补译",
                f"所在景点据{zone[2]}原文：「{zone[1]}」" if zone else "", note_en) if x)))
    return rows


# ---------------------------------------------------------------- 序号
def assign_seq(rows: list[dict], write: bool, r) -> None:
    old: dict[str, int] = {}
    if SEQ_MAP.exists():
        with open(SEQ_MAP, encoding="utf-8", newline="") as fh:
            for row in csv.DictReader(fh):
                old[row["key"]] = int(row["seq"])
    if len(set(old.values())) != len(old):
        sys.exit(f"[fatal] {SEQ_MAP.name} 里有重复的 seq")
    nxt = max(old.values(), default=0) + 1
    new_keys = []
    for row in rows:
        if row["key"] not in old:
            old[row["key"]] = nxt
            new_keys.append(row["key"])
            nxt += 1
        row["seq"] = old[row["key"]]
    keys = [row["key"] for row in rows]
    if len(set(keys)) != len(keys):
        sys.exit(f"[fatal] 身份键重复：{[k for k, n in Counter(keys).items() if n > 1][:10]}")
    gone = sorted(set(old) - set(keys), key=lambda k: old[k])
    r.say(f"\n## 序号：沿用 {len(rows) - len(new_keys)}，新增 {len(new_keys)}，"
          f"本次不再出现（留空号，不回收）{len(gone)} {gone[:8]}")
    if write:
        with open(SEQ_MAP, "w", encoding="utf-8", newline="") as fh:
            w = csv.writer(fh)
            w.writerow(["key", "seq"])
            for k, s in sorted(old.items(), key=lambda x: x[1]):
                w.writerow([k, s])


# ---------------------------------------------------------------- 体检
def health(rows: list[dict], r) -> None:
    r.say("\n## 体检")
    zh = [x for x in rows if x["key"].startswith("zh:")]
    r.say(f"总 {len(rows)} 行：中文站 {len(zh)}（其中有馆方英文 {sum(1 for x in zh if x['name_en'])}），"
          f"仅英文站 {len(rows) - len(zh)}")
    L = [len(x["desc"]) for x in zh]
    r.say(f"中文简介：中位 {statistics.median(L)} 字，最短 {min(L)}，最长 {max(L)}；有图 {sum(1 for x in rows if x['image'])}/{len(rows)}")
    E = [len(x["desc_en"]) for x in rows if x["desc_en"]]
    r.say(f"馆方英文简介 {len(E)} 条：中位 {statistics.median(E)} 字符，最短 {min(E)}，最长 {max(E)}")
    r.say(f"开放状态 {dict(Counter(x['on'] if x['on'] == ON_NO else '向游人开放（原文写明）' for x in rows))}")
    r.say(f"所在景点 {dict(Counter(x['hall'] or '（空）' for x in rows))}")
    for x in sorted(rows, key=lambda x: x["seq"]):
        r.say(f"  {x['seq']:>3} {x['name'] or '—':<8} | {x['name_en'] or '—':<36} | {x['hall'] or '—':<4} | "
              f"中 {len(x['desc']):>4} 字 英 {len(x['desc_en']):>4} | {x['on'][:30]}")
    for x in rows:
        if x["note"]:
            r.say(f"  备注 {x['seq']:>3} {x['name'] or x['name_en']}：{x['note']}")


# ---------------------------------------------------------------- 写 Excel
HEADER = ["序号", "展厅", "展品名称", "展品简介", "展品图片", "陈列状态", "官方页面", "Tier",
          "对象层级", "类目", "英文名称", "英文简介", "官网ID", "出处", "英文页面", "备注"]


def write_xlsx(rows: list[dict], d: dict, r) -> None:
    import openpyxl
    from openpyxl.styles import Alignment, Font
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "展品清单"
    ws.append(["北海公园 景点清单"])
    ws.append([f"资料整理日期：{dt.date.today():%Y 年 %m 月} | 数据来源：{BASE} 官方接口（抓取 {d['fetched']}） | "
               f"景点 {len(rows)} 个 | 由 beihai_build.py 生成，勿手工编辑"])
    ws.append([])
    ws.append(HEADER)
    for c in ws[4]:
        c.font = Font(bold=True)
    for x in sorted(rows, key=lambda x: x["seq"]):
        ws.append([x["seq"], x["hall"], x["name"], x["desc"], x["image"], x["on"], x["url"], "",
                   x["level"], x["cat"], x["name_en"], x["desc_en"], x["oid"], x["src"], x["en_url"], x["note"]])
    for col, width in zip("ABCDEFGHIJKLMNOP", (6, 8, 14, 70, 30, 24, 30, 6, 8, 14, 22, 60, 26, 22, 30, 50)):
        ws.column_dimensions[col].width = width
    for row in ws.iter_rows(min_row=5):
        for c in row:
            c.alignment = Alignment(wrap_text=True, vertical="top")

    kf = json.loads(d["trees"]["guide"]["guide-kfsj"]["columnInfo"]["fdExtendInfo"])
    info = wb.create_sheet("参观实用信息")
    info.append(["北海公园 — 参观实用信息"])
    info.append([])
    info.append(["项目", "内容"])
    for k, v in (
        ("开放时间", f"旺季（{kf['peakDate1']}）{kf['kfsjPeakOpenTime']}，{kf['peakLastEntryTime']} 停止入园；"
                     f"淡季（{kf['offDate1']}）{kf['kfsjOffOpenTime']}，{kf['offLastEntryTime']} 停止入园。"
                     f"园中园及小庭院：旺季 {kf['peakGardenOpenTime']}–{kf['peakGardenCloseTime']}（{kf['peakGardenLastEntryTime']} 停止入园），"
                     f"淡季 {kf['offGardenOpenTime']}–{kf['offGardenCloseTime']}（{kf['offGardenLastEntryTime']} 停止入园）"),
        ("闭园日", "游园须知：「园中园景点（团城除外）逢周一关闭，法定节假日除外」。官网没有列出哪些景点算园中园"),
        ("票价", "旺季门票 10 元、联票 20 元；淡季门票 5 元、联票 15 元（游园须知）。官网没有列出联票含哪些景点"),
        ("数据范围", "官网「景点介绍」中文 21 条 + 英文站独有的琼华岛、白塔 2 条，共 23 个景点"),
        ("数据来源", f"{BASE} 的 /api 接口（{d['fetched']} 抓取，beihai_site_scrape.py）。境外网络连不上，须在国内网络抓"),
        ("开放状态", "官网没有逐个写明是否开放。原文写明「向游人开放」的记为开放，其余记「开放状态未知」"),
        ("英文", "英文站有的 8 条用馆方英文；其余由 AI 补译。中文站标题后附带的英文是机器直译，未采用"),
        ("图片版权", "图片地址指向北海公园官网服务器，版权归原单位，仅供个人查阅"),
    ):
        info.append([k, v])
    info.column_dimensions["A"].width = 12
    info.column_dimensions["B"].width = 110
    XLSX.parent.mkdir(exist_ok=True)
    wb.save(XLSX)
    r.say(f"\n写出 {XLSX.relative_to(HERE)}（{len(rows)} 行）")


class Report:
    def __init__(self):
        self.lines: list[str] = []

    def say(self, *a) -> None:
        line = " ".join(str(x) for x in a)
        self.lines.append(line)
        print(line, flush=True)


def main() -> None:
    sys.stdout.reconfigure(encoding="utf-8")        # 中文 Windows 控制台默认 GBK；
    sys.stderr.reconfigure(encoding="utf-8")        # sys.exit 的报错走 stderr，也要改
    ap = argparse.ArgumentParser(description=__doc__.split("\n\n")[0])
    ap.add_argument("--check", action="store_true", help="只做配对与体检，不写文件")
    args = ap.parse_args()
    r = Report()
    r.say(f"# 北海公园 源表生成{'（只检查）' if args.check else ''}")
    d = load()
    rows = build_rows(d, r)
    assign_seq(rows, write=not args.check, r=r)
    health(rows, r)
    if not args.check:
        write_xlsx(rows, d, r)
        REPORT.write_text("\n".join(r.lines) + "\n", "utf-8")


if __name__ == "__main__":
    main()
