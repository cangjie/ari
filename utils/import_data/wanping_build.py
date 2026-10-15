#!/usr/bin/env python3
"""宛平城主要参观点：由已保存的官方网页摘录确定性生成清单，零网络、零模型。

范围由用户 2026-10-15 确认：宛平城为主，含卢沟桥、抗战馆、雕塑园主要参观点。
输入 wanping_source_capture.json 是网页工具提取的文本，不冒充原始 HTTP 响应。
正文逐字摘录；编者说明单列；未取得正文的来源不得作为已完成条目。
不写数据库，不改任何存量馆，不生成路线或无出处的坐标。
"""
from __future__ import annotations

import argparse
import csv
import hashlib
import json
import re
from pathlib import Path

HERE = Path(__file__).resolve().parent
CAPTURE = HERE / 'wanping_source_capture.json'
SEQ = HERE / 'wanping_seq_map.csv'
DATA = HERE / 'wanping_site_data.json'
XLSX = HERE / 'artworks' / '宛平城_参观点清单.xlsx'

# 身份键固定；名称是目录标签，不冒充文物编目题名。每个标签都有逐字锚点。
# key, name, parent, type hint, [(source id, first line, last line)], note
CANDIDATES = [
 ('city','宛平城','', 'site',[('wanping_overview',105,107),('wanping_rebuilt',82,88)],'两份官方原文均保留；纪年与旧城名有待审计，不自行改写。'),
 ('east_gate','顺治门','city','node',[('gate_history',37,38)],'宛平城东门；本条为古城组成部分。'),
 ('west_gate','永昌门（威严门）','city','node',[('gate_history',37,38)],'宛平城西门；永昌门与威严门是同一门的不同时期名称，不拆两条。'),
 ('walls','宛平城墙','city','node',[('area_renewal_2025',114,120)],'文物本体修缮与周边院落复建分开理解；不是所有城内建筑都属于明代原构。'),
 ('shell_scars','城墙弹痕','walls','node',[('wanping_rebuilt',86,86),('area_renewal_2025',132,140)],'城墙组成部分；尚无逐处弹坑定位与编号，不编坐标。'),
 ('wall_path','环城步道','city','node',[('area_renewal_2025',99,100),('area_open_2025',129,129)],'2025年报道的城外步道；不同于城墙顶部登城游览。'),
 ('main_street','宛平大街','city','node',[('wanping_rebuilt',88,88)],'官方将已复建街区称为明清文化步行街，不等同于整条街保存了明清原构。'),
 ('county_office','宛平县衙','city','node',[('wanping_rebuilt',88,88),('gate_history',51,52)],'官方明确记为复建；与历史县署驻地的沿革分开记录。当前室内开放状态未核实。'),
 ('post_station','卢沟驿','city','node',[('wanping_rebuilt',88,88)],'官方明确记为复建；当前内部展陈与室内开放状态未核实。'),
 ('xinglong_temple','兴隆寺','city','node',[('wanping_rebuilt',88,88),('area_renewal_2025',127,127)],'官方明确记为复建；2025年报道记为兼具休憩与研学功能的公共空间。'),
 ('gongji_camp','拱极营','city','node',[('wanping_rebuilt',88,88),('area_renewal_2025',104,104)],'复建建筑内的冷兵器展馆；未逐件收集刀剑等藏品，不能把所有陈列都称为出土原物。'),
 ('theatre','戏楼','city','node',[('wanping_rebuilt',88,88)],'官方称为增加的景观；未提供具体建筑沿革和当前演出安排。'),
 ('city_exhibition','烽火卢沟 硬气铸魂','city','node',[('area_renewal_2025',107,107)],'位于西门内小城会客厅的丰台抗战专题展；本条与场馆合为一个参观节点。2025年报道不证明今天仍在展。'),
 ('bridge','卢沟桥','', 'node',[('bridge_overview',105,107),('bridge_tourism',41,41)],'桥宽口径、石狮数量保留各来源原文；不同年份的数字不相互覆盖。'),
 ('bridge_lions','卢沟桥石狮','bridge','object',[('bridge_overview',105,105),('bridge_tourism',48,51)],'作为石狮群记录，不把未取得单独身份资料的数百只石狮分别编目；官方来源数量口径不一致。'),
 ('moon_stele','卢沟晓月碑','bridge','object',[('bridge_tourism',55,55)],'目录名称据原文“卢沟晓月”碑；原物与碑亭不拆成两件。'),
 ('kangxi_stele','康熙重修卢沟桥御制碑','bridge','object',[('bridge_tourism',55,55)],'目录名称据原文“清康熙帝于1698年为记述重修卢沟桥而竖的御制碑”；未取得碑文全文。'),
 ('huabiao','卢沟桥华表','bridge','object',[('bridge_tourism',51,51)],'桥两头华表作为组合记录；未虚构逐件文物编号。'),
 ('war_museum','中国人民抗日战争纪念馆','', 'site',[('museum_overview',105,105),('area_renewal_2025',146,149)],'采用2025年改扩建后的基本陈列，不沿用2015年旧展厅目录。'),
 ('permanent_exhibition','为了民族解放与世界和平','war_museum','node',[('museum_exhibition_2025',4,8)],'完整名称为《为了民族解放与世界和平——纪念中国人民抗日战争暨世界反法西斯战争胜利80周年主题展览》；馆方明确作为基本陈列长期展出。'),
 ('exhibit_1','奋起抵抗 中国人民打响世界反法西斯战争第一枪','permanent_exhibition','node',[('museum_exhibition_2025',16,16)],'基本陈列第一部分；未取得楼层和展厅编号。'),
 ('exhibit_2','全民族抗战 开辟世界反法西斯战争东方主战场','permanent_exhibition','node',[('museum_exhibition_2025',27,27)],'基本陈列第二部分；未取得楼层和展厅编号。'),
 ('exhibit_3','砥柱中流 中国共产党坚定不移站在抗战最前线','permanent_exhibition','node',[('museum_exhibition_2025',37,37)],'基本陈列第三部分；未取得楼层和展厅编号。'),
 ('exhibit_4','战略支撑 中国艰苦抗战打破法西斯称霸世界的图谋','permanent_exhibition','node',[('museum_exhibition_2025',47,47)],'基本陈列第四部分；未取得楼层和展厅编号。'),
 ('exhibit_5','罄竹难书 日本侵略者的法西斯暴行','permanent_exhibition','node',[('museum_exhibition_2025',57,57)],'基本陈列第五部分；未取得楼层和展厅编号。'),
 ('exhibit_6','得道多助 世界各国人民对中国抗战的支持','permanent_exhibition','node',[('museum_exhibition_2025',66,66)],'基本陈列第六部分；未取得楼层和展厅编号。'),
 ('exhibit_7','伟大胜利 中华民族走向复兴的历史转折点','permanent_exhibition','node',[('museum_exhibition_2025',76,76)],'基本陈列第七部分；未取得楼层和展厅编号。'),
 ('exhibit_8','铭记历史 珍爱和平 携手构建人类命运共同体','permanent_exhibition','node',[('museum_exhibition_2025',86,86)],'基本陈列第八部分；未取得楼层和展厅编号。'),
 ('heroes_zone','弘扬伟大抗战精神 深切缅怀抗战英烈','permanent_exhibition','node',[('museum_exhibition_2025',90,90)],'基本陈列中的专区；不把家书等未独立编目的实物另列藏品。'),
 ('awakening_lion','东方醒狮','city','object',[('area_open_2025',113,113)],'位于抗战馆前广场南侧；未取得雕塑尺寸和作者的可靠资料。'),
 ('sculpture_park','中国人民抗日战争纪念雕塑园','', 'site',[('sculpture_park_overview',105,105),('area_open_2025',129,129)],'与宛平城、卢沟桥、抗战馆有包含和相邻关系；清单不是路线，不把父子节点停留时间相加。'),
 ('war_monument','中国人民抗日战争纪念碑','sculpture_park','object',[('sculpture_park_overview',105,105)],'雕塑园中的纪念碑；尺寸原文照录，不推广到其他雕塑。'),
 ('park_1','日寇侵凌','sculpture_park','node',[('area_open_2025',129,129)],'雕塑群四大主题分区之一；未取得本区逐件雕塑清单及边界。'),
 ('park_2','奋起救亡','sculpture_park','node',[('area_open_2025',129,129)],'雕塑群四大主题分区之一；未取得本区逐件雕塑清单及边界。'),
 ('park_3','抗日烽火','sculpture_park','node',[('area_open_2025',129,129)],'雕塑群四大主题分区之一；未取得本区逐件雕塑清单及边界。'),
 ('park_4','正义必胜','sculpture_park','node',[('area_open_2025',129,129)],'雕塑群四大主题分区之一；未取得本区逐件雕塑清单及边界。'),
 ('living_wall','血肉长城','sculpture_park','object',[('area_open_2025',129,129)],'官方2025年报道点名的雕塑；本次不推断它属于哪个主题分区。'),
 ('field_rescue','战地救援','sculpture_park','object',[('area_renewal_2025',162,166)],'官方2025年报道点名的雕塑；原文中的医生形象阐释照录，不把阐释改写为作者确认的肖像身份。'),
 ('bridge_museum','卢沟桥历史博物馆','bridge','node',[('bridge_museum_index',0,0)],'仅取得北京旅游网的官方搜索索引摘录；详情正文超时。未据此采纳实时票价、开放时间或单件文物位置。'),
]

def main():
    ap = argparse.ArgumentParser(description=__doc__)
    ap.add_argument('--check', action='store_true')
    args = ap.parse_args()
    captured = json.loads(CAPTURE.read_text('utf-8'))
    sources = {s['id']: s for s in captured['sources']}
    indexed=json.loads((HERE/'wanping_bridge_museum_index.json').read_text('utf-8'))
    sources[indexed['source_id']]={**indexed,'capture':'L0: '+indexed['quote']}
    # 网页工具会把相邻的图片/空白行压在同一行；先恢复每个 L<n>: 标记的边界。
    lines = {k: {int(n): t for n,t in re.findall(r'^L(\d+): ?(.*)$',
              re.sub(r'(?<!\n)(L\d+:)',r'\n\1',s['capture']),re.M)} for k,s in sources.items()}
    old = {}
    if SEQ.exists():
        with SEQ.open(encoding='utf-8',newline='') as f:
            for r in csv.DictReader(f):
                if r['identity_key'] in old or int(r['source_seq']) in old.values():
                    raise ValueError('序号映射重复')
                old[r['identity_key']] = int(r['source_seq'])
    seq_map = dict(old)
    next_seq = max(seq_map.values(),default=0)+1
    rows=[]
    for key,name,parent,kind,refs,note in CANDIDATES:
        if key not in seq_map:
            seq_map[key]=next_seq;next_seq+=1
        evidence=[]
        for sid,a,b in refs:
            content=lines[sid]
            if any(n not in content for n in range(a,b+1)):
                raise ValueError(f'{sid}: 未完整保存所需行 {a}-{b}')
            quote='\n'.join(content[n] for n in range(a,b+1) if content[n] and not content[n].startswith('cite'))
            if not quote or 'Failed to fetch' in quote:
                raise ValueError(f'{key}: 正文不可用')
            evidence.append(dict(source_id=sid,url=sources[sid]['url'],lines=[a,b],quote=quote,
                                 retrieval_method=sources[sid].get('retrieval_method','web_page_extracted_text')))
        rows.append(dict(source_seq=seq_map[key],identity_key=key,name=name,parent_key=parent,
                         object_type_hint=kind,description='\n\n'.join(e['quote'] for e in evidence),
                         editor_note=note,status='未知（资料记载时间不等于实时开放状态）',evidence=evidence,
                         image_url=None,coordinates=None))
    assert len(rows)==len({r['identity_key'] for r in rows})==39
    assert all(not r['parent_key'] or r['parent_key'] in seq_map for r in rows)
    result=dict(museum_key='wanping',name='宛平城及周边主要参观点',collected_on=captured['collected_on'],
                scope=captured['scope'],capture_sha256=hashlib.sha256(CAPTURE.read_bytes()).hexdigest(),
                models={'information':'gpt-5.6-luna','rating':'gpt-5.6-luna high/medium/xhigh (codex-cli)',
                        'audit':'claude-sonnet-5 effort=medium (claude-cli), slim'},
                state='来源清单；评分与审计尚未完成；未入库',rows=rows,
                pending=[{'name':'卢沟桥历史博物馆','url':'https://s.visitbeijing.com.cn/attraction/120799',
                          'reason':'已按官方索引摘录收为节点，详情正文仍超时；不是完整抓取，开放安排和展品需补核。'},
                         {'name':'武俊碑、武俊腰牌','reason':'官方报道有介绍，但本轮未核实当前展陈位置及原件/复制品，暂不列为可参观单件。'}])
    print(json.dumps({'records':len(rows),'types':{k:sum(r['object_type_hint']==k for r in rows) for k in ['site','node','object']},
                      'captured_sources':len(sources),'pending':len(result['pending'])},ensure_ascii=False))
    if args.check:return
    DATA.write_text(json.dumps(result,ensure_ascii=False,indent=2)+'\n',encoding='utf-8')
    with SEQ.open('w',encoding='utf-8',newline='') as f:
        w=csv.writer(f);w.writerow(['identity_key','source_seq']);w.writerows(sorted(seq_map.items(),key=lambda x:x[1]))
    write_xlsx(result)

def write_xlsx(data):
    from openpyxl import Workbook
    from openpyxl.styles import Font, PatternFill, Alignment
    from openpyxl.utils import get_column_letter
    wb=Workbook();ws=wb.active;ws.title='参观点清单'
    ws.append(['序号','父节点','参观点名称','官方原文摘录','图片','状态','官方页面','Tier','对象层级提示','身份键','编者说明'])
    for r in data['rows']:
        ws.append([r['source_seq'],r['parent_key'],r['name'],r['description'],None,r['status'],r['evidence'][0]['url'],None,r['object_type_hint'],r['identity_key'],r['editor_note']])
    ws.freeze_panes='D2';ws.auto_filter.ref=ws.dimensions
    for i,width in enumerate([8,22,48,100,16,40,65,9,16,22,85],1):ws.column_dimensions[get_column_letter(i)].width=width
    for row in ws:
        for c in row:c.font=Font(name='Arial',size=11);c.alignment=Alignment(vertical='top',wrap_text=True)
    for c in ws[1]:c.fill=PatternFill('solid',fgColor='234E52');c.font=Font(name='Arial',bold=True,color='FFFFFF')
    for n in range(2,ws.max_row+1):ws.row_dimensions[n].height=110
    ev=wb.create_sheet('出处逐条');ev.append(['序号','参观点名称','来源标识','URL','已保存行号','逐字摘录'])
    for r in data['rows']:
        for e in r['evidence']:ev.append([r['source_seq'],r['name'],e['source_id'],e['url'],str(e['lines']),e['quote']])
    ev.freeze_panes='C2';ev.auto_filter.ref=ev.dimensions
    for col,width in zip('ABCDEF',[8,48,27,65,18,110]):ev.column_dimensions[col].width=width
    info=wb.create_sheet('说明与待补');info.append(['项目','说明'])
    for k,v in [('范围',data['scope']),('采集日期',data['collected_on']),('状态',data['state']),('来源性质','网页工具提取文本，非原始HTTP响应；官方原文与编者说明分列。'),('层级','父节点与子节点有包含关系，本表不是互不重叠的路线站点。'),('开放状态','历史报道或长期陈列安排不等于今天开放；票价、开放时间与室内位置均需临行复核。')]:info.append([k,v])
    for p in data['pending']:info.append([p['name'],p['reason']])
    info.column_dimensions['A'].width=30;info.column_dimensions['B'].width=120
    for sheet in [ev,info]:
        for row in sheet:
            for c in row:c.font=Font(name='Arial',size=11);c.alignment=Alignment(vertical='top',wrap_text=True)
        for c in sheet[1]:c.font=Font(name='Arial',bold=True,color='FFFFFF');c.fill=PatternFill('solid',fgColor='234E52')
    wb.save(XLSX)
    print(XLSX)

if __name__=='__main__':main()
