#!/usr/bin/env python3
# -*- coding: utf-8 -*-
"""
从抓取数据确定性地生成吉林省博物院的源 Excel，零网络、零 API。

    python3 jlpm_build.py            # 写 artworks/吉林省博物院_展品清单.xlsx + jlpm_seq_map.csv + 体检报告
    python3 jlpm_build.py --check    # 只做合并、名录核对与体检，不写文件

输入（都已提交）：
    jlpm_site_data.py            镇馆之宝 17 件、展览全列（jlpm_site_scrape.py --scrape）
    jlpm_collectdb.jsonl         藏品数据库 17709 件（同上）
    jlpm_wikidata_catalog.json   Wikidata 上的《全国馆藏文物名录》11947 条（jlpm_wikidata.py）

**范围是用户逐项定的**（2026-09-26/27），不在这里自行扩大：
· 藏品：镇馆之宝与藏品数据库**全部入库**，**只评有介绍的**（「评分范围」列）。
  藏品数据库 17709 件里 17514 件介绍为空、105 件介绍只是重复名称，有实质介绍的约 90 件。
· 节点：当前在展的实体展览（馆方标记「展出中」/「即将展出」），另加「白山松水的记忆」
  （标记「已结束」、时间栏却写「正在展出」—— 收为一个节点，在展记「未核实」）。
· 名录：只用来核对身份，名称**逐字相同且两边都唯一**才挂可移动文物编号，不做模糊匹配。

**在展状态**：官网给不出「哪件东西在哪个展」（展品清单接口全部 0 件；镇馆之宝详情的
「曾经展出」对每件都是同一串最新展览，抓取时已剔除）。所以藏品一律「在展状态未知」、展厅留空。
出现在「镇馆之宝」里 ≠ 在展（MFA「来源是官网 ≠ 在展」那条教训）。

**序号按身份键幂等**（AGENTS.md 第 14 条）：身份键 -> source_seq 存在 jlpm_seq_map.csv，
已有的键永不改号，新键追加在末尾，消失的键留空号。首次编号时评分范围排在前面
（镇馆之宝、有介绍的数据库藏品、节点），只有名称的排在后面 —— 只影响首次，不影响幂等。
"""

from __future__ import annotations

import argparse
import csv
import datetime as dt
import json
import pathlib
import re
import statistics
import sys
from collections import Counter

import jlpm_site_data as site

HERE = pathlib.Path(__file__).resolve().parent
BASE = "https://www.jlmuseum.net"
XLSX = HERE / "artworks" / "吉林省博物院_展品清单.xlsx"
SEQ_MAP = HERE / "jlpm_seq_map.csv"
DB_FILE = HERE / "jlpm_collectdb.jsonl"
CATALOG = HERE / "jlpm_wikidata_catalog.json"
REPORT = HERE / "jlpm_samples" / "build_report.txt"

SCOPE_YES = "评分"
SCOPE_NO = "不评分（只有名称）"

# ---------------------------------------------------------------- 镇馆之宝 ↔ 藏品数据库
# 同一件东西在两个栏目里各有一条：镇馆之宝用精品栏目的名称（不带朝代），数据库用名录写法
# （带朝代前缀）。**逐对核实**，每条写明凭什么；按名称包含去配会配错 —— 「契丹文铜镜」
# 在库里同时能包含进「辽契丹文铜镜」（只有名称）与「辽契丹文八角铜镜」，真正对上的是后者。
# 合并后一行：名称用镇馆之宝的，数据库名称进「别名」（名录核对用它），两段介绍都保留。
TWIN = {
    "5":  ("元青花云龙纹高足碗", "去掉朝代前缀后名称逐字相同、朝代同为元，库中唯一；库里无尺寸与介绍可比"),
    "7":  ("东汉鎏金神兽铜牌饰", "尺寸 11、7.2、0.15 三个数一致，朝代同为东汉"),
    "12": ("汉白玉耳杯", "尺寸 3.2、13、9.5 三个数一致"),
    "4":  ("辽石雕彩绘塔", "尺寸 96.5、45、37、10 四个数一致"),
    "11": ("辽契丹文八角铜镜", "两段介绍都写 1971 年大安县红岗出土、八角形、契丹小字铭文五行、陈述先生释读；"
                              "直径 26.2 / 26。**不是**库里同样能包含进来的「辽契丹文铜镜」（只有名称）"),
    "13": ("明“禾屯吉卫指挥使司印” 铜印", "铭文、印背「礼部造」「永乐七年九月日」、侧边「礼字四十三号」、1974 年出土全同；"
                                      "出土地一处写洮安县、一处写洮南市（同一地的旧名与新名），边长 9 / 8.8"),
    "10": ("渤海石狮", "都是 1949 年敦化贞惠公主墓出土，通高 64 厘米"),
    "9":  ("汉错金银“丙午神钩”铜带钩", "长 15.7 一致，名称除朝代前缀外相同（镇馆之宝标东汉，库名写汉）"),
    "20": ("元龙泉窑观音像", "高 23.8 一致，朝代同为元"),
    "21": ("辽银釉鸡冠壶", "尺寸 25、11.5、5.6、10 四个数一致"),
    "24": ("清郑燮竹石图轴", "纵 130.5 一致，横 71.5 / 71.3；都是郑燮（板桥）"),
}
# 找到过候选但**不合并**的，只进报告：证据不足，宁可两行各算一件
TWIN_REJECTED = {
    "8": "定窑紫釉印花碗（金）↔ 库里「宋定窑紫釉印花瓷碗」有多件、朝代不同",
    "2": "龙泉窑豆青釉琮式瓶（辽，高 18）↔ 库里「宋龙泉窑青釉琮式瓷瓶」（宋，高 18.5）：朝代不同",
}

# ---------------------------------------------------------------- 节点
FLAG_ON = {"1": "展出中（馆方标记）", "2": "即将展出（馆方标记）"}
BSSS = "299"          # 白山松水的记忆
ON_BSSS = "在展状态未核实（馆方标记「已结束」，时间栏写「正在展出」）"
ON_UNKNOWN = "馆藏（在展状态未知）"

# 展览的「地点」写法五花八门，这里逐条归到展厅名。**没登记的地点直接报错**，不猜
PLACE_HALL = {
    "吉林省近现代史展（春京西）": "吉林省近现代史展（春京西）",
    "吉林省博物院3楼B区展厅": "3楼B区",
    "3楼A区": "3楼A区",
    "1楼A区（吉林省博物院）": "1楼A区",
    "1楼B区（吉林省博物院）": "1楼B区",
    "1楼C区（吉林省博物院）": "1楼C区",
    "1楼环廊（吉林省博物院）": "1楼环廊",
    "吉林省博物院二楼": "2楼",
}


def _txt(s) -> str:
    """HTML/带换行的原文 -> 一段纯文本（段落间留一个空格）。"""
    s = re.sub(r"<br\s*/?>|</p>", "\n", str(s or ""), flags=re.I)
    s = re.sub(r"<[^>]+>", "", s).replace("&nbsp;", " ")
    return re.sub(r"[ \t　]*\n[\s　]*", "\n", s).strip()


def _flat(s: str) -> str:
    return re.sub(r"\s+", "", s or "")


def _substantive(content: str, name: str) -> bool:
    """有实质介绍：去空白后至少 15 字，且比名称多出 4 字以上（很多条的介绍只是把名称重抄一遍）。"""
    c = _flat(_txt(content))
    return len(c) >= 15 and len(c) > len(_flat(name)) + 4


def _abs(u: str) -> str:
    return BASE + u if u and u.startswith("/") else (u or "")


def load_db() -> list[dict]:
    with open(DB_FILE, encoding="utf-8") as fh:
        return [json.loads(line) for line in fh if line.strip()]


def build_rows(r) -> list[dict]:
    db = load_db()
    by_name: dict[str, list[dict]] = {}
    for x in db:
        by_name.setdefault(x["name"], []).append(x)
    twin_db: dict[str, dict] = {}
    for tid, (dname, why) in TWIN.items():
        hit = by_name.get(dname, [])
        if len(hit) != 1:
            sys.exit(f"[fatal] TWIN {tid} 登记的库名「{dname}」在藏品数据库里有 {len(hit)} 条（应为 1）")
        twin_db[tid] = hit[0]
    twin_ids = {x["id"] for x in twin_db.values()}

    rows: list[dict] = []
    for tid, t in sorted(site.TREASURE.items(), key=lambda x: int(x[0])):
        d = twin_db.get(tid)
        desc = _txt(t.get("contentNoHtml") or t.get("content"))
        if d and _substantive(d.get("content", ""), d["name"]) and _flat(_txt(d["content"])) not in _flat(desc):
            desc += "\n【藏品数据库介绍】" + _txt(d["content"])
        rows.append(dict(
            key=f"t:{tid}", hall="", name=t["name"].strip(), desc=desc,
            image=_abs(t.get("mainImgUrl") or ""), on=ON_UNKNOWN,
            url=f"{BASE}/#/collect/detail?id={tid}", level="藏品", cat=t.get("typeName") or "",
            oid=f"collect/{tid}" + (f" · collectdb/{d['id']}" if d else ""),
            src="jlpm_official（镇馆之宝）" + (" + jlpm_collectdb" if d else ""),
            scope=SCOPE_YES, year=t.get("yearTypeName") or "", texture=t.get("texture") or "",
            size=(t.get("size") or "").strip(), grade=t.get("levelName") or "",
            alias=d["name"] if d else "", twin_why=TWIN.get(tid, ("", ""))[1]))

    rest = []
    for x in db:
        if x["id"] in twin_ids:
            continue
        rich = x.get("detail_ok") and _substantive(x.get("content", ""), x["name"])
        img = x.get("mainImgUrl") or (x.get("imgList") or [""])[0]
        row = dict(
            key=f"d:{x['id']}", hall="", name=x["name"].strip(),
            desc=_txt(x.get("content")) if rich else "", image=img, on=ON_UNKNOWN,
            url=f"{BASE}/#/collect/collectionDatabase?id={x['id']}", level="藏品",
            cat=x.get("typeName") or "", oid=f"collectdb/{x['id']}",
            src="jlpm_collectdb" + ("" if x.get("detail_ok") else "（详情取不到，只有列表信息）"),
            scope=SCOPE_YES if rich else SCOPE_NO, year=x.get("yearTypeName") or "",
            texture=x.get("texture") or "", size=(x.get("size") or "").strip(), grade="",
            # 21 件吴大澂作品列表上是具体题名、详情里是名录通称，两者都是馆方的，通称进别名
            alias=(x.get("detail_name") or "").strip(), twin_why="")
        (rows if rich else rest).append(row)

    for k, e in sorted(site.EXHIBITIONS.items(), key=lambda x: int(x[0])):
        flag = str(e.get("exhibitionFlag"))
        if e.get("lists") == ["虚拟展厅"] or not (flag in FLAG_ON or k == BSSS):
            continue
        place = (e.get("place") or "").strip()
        if place not in PLACE_HALL:
            sys.exit(f"[fatal] 展览 {k}「{e['name']}」的地点「{place}」没在 PLACE_HALL 里登记")
        basic = "基本陈列" in e.get("lists", [])
        rows.append(dict(
            key=f"x:{k}", hall=PLACE_HALL[place], name=e["name"].strip(), desc=_txt(e.get("content")),
            image=_abs(e.get("showPic") or ""), on=ON_BSSS if k == BSSS else FLAG_ON[flag],
            url=f"{BASE}/#/exhibition/detail?id={k}", level="基本陈列" if basic else "临时展览",
            cat="展览", oid=f"exhibition/{k}", src="jlpm_official（展览）", scope=SCOPE_YES,
            year="", texture="", size="", grade="", alias="",
            twin_why=f"地点原文：{place}；展期：{(e.get('exhibitionTime') or '').strip()}"))
    return rows + rest


# ---------------------------------------------------------------- 名录核对（身份）
def match_catalog(rows: list[dict], r) -> None:
    """名称（及别名）与名录逐字相同（只去空白），且在本馆各行里唯一、在名录里唯一，才挂编号。
    一行的两个名字若指向名录里不同的两条，一概不挂。"""
    cat = json.loads(CATALOG.read_text("utf-8"))["items"]
    cat_by: dict[str, list[dict]] = {}
    for c in cat:
        cat_by.setdefault(_flat(c["name"]), []).append(c)
    ours = Counter(_flat(n) for x in rows if x["level"] == "藏品" for n in {x["name"], x["alias"]} if n)
    used: Counter = Counter()
    for x in rows:
        x["relic"], x["qid"] = "", ""
        if x["level"] != "藏品":
            continue
        hits = {c["qid"]: c for n in {x["name"], x["alias"]} if n
                for c in cat_by.get(_flat(n), []) if len(cat_by[_flat(n)]) == 1 and ours[_flat(n)] == 1}
        if len(hits) == 1:
            c = next(iter(hits.values()))
            x["qid"], x["relic"] = c["qid"], ";".join(c["relic_id"])
            used[c["qid"]] += 1
    dup = {q for q, n in used.items() if n > 1}
    for x in rows:
        if x.get("qid") in dup:           # 名录里同一条挂到了两行：两行都不挂
            x["qid"] = x["relic"] = ""
    objs = [x for x in rows if x["level"] == "藏品"]
    sc = [x for x in objs if x["scope"] == SCOPE_YES]
    r.say(f"\n## 名录核对（逐字相同且两边唯一）")
    r.say(f"  藏品 {len(objs)} 件挂上编号 {sum(1 for x in objs if x['qid'])}；"
          f"其中评分范围 {len(sc)} 件挂上 {sum(1 for x in sc if x['qid'])}；同一条挂到两行而撤回的 {len(dup)}")
    r.say(f"  评分范围里没挂上的：{[x['name'] for x in sc if not x['qid']][:60]}")


# ---------------------------------------------------------------- 序号
def assign_seq(rows: list[dict], write: bool, r) -> None:
    old: dict[str, int] = {}
    if SEQ_MAP.exists():
        with open(SEQ_MAP, encoding="utf-8", newline="") as fh:
            for row in csv.DictReader(fh):
                old[row["key"]] = int(row["seq"])
    if len(set(old.values())) != len(old):
        sys.exit(f"[fatal] {SEQ_MAP.name} 里有重复的 seq")
    nxt = max(old.values(), default=0) + 1
    new_keys = []
    for row in rows:
        if row["key"] not in old:
            old[row["key"]] = nxt
            new_keys.append(row["key"])
            nxt += 1
        row["seq"] = old[row["key"]]
    keys = [row["key"] for row in rows]
    if len(set(keys)) != len(keys):
        sys.exit(f"[fatal] 身份键重复：{[k for k, n in Counter(keys).items() if n > 1][:10]}")
    gone = sorted(set(old) - set(keys), key=lambda k: old[k])
    r.say(f"\n## 序号：沿用 {len(rows) - len(new_keys)}，新增 {len(new_keys)}，"
          f"本次不再出现（留空号，不回收）{len(gone)} {gone[:8]}")
    if write:
        with open(SEQ_MAP, "w", encoding="utf-8", newline="") as fh:
            w = csv.writer(fh)
            w.writerow(["key", "seq"])
            for k, s in sorted(old.items(), key=lambda x: x[1]):
                w.writerow([k, s])


# ---------------------------------------------------------------- 体检
def health(rows: list[dict], r) -> None:
    r.say("\n## 体检")
    objs = [x for x in rows if x["level"] == "藏品"]
    nodes = [x for x in rows if x["level"] != "藏品"]
    sc = [x for x in rows if x["scope"] == SCOPE_YES]
    r.say(f"总 {len(rows)} 行：藏品 {len(objs)}（镇馆之宝 {sum(1 for x in objs if x['key'].startswith('t:'))}，"
          f"其中与数据库合并 {len(TWIN)} 对），节点 {len(nodes)} {dict(Counter(x['level'] for x in nodes))}")
    r.say(f"评分范围 {len(sc)} 行：镇馆之宝 {sum(1 for x in sc if x['key'].startswith('t:'))}、"
          f"数据库有介绍的 {sum(1 for x in sc if x['key'].startswith('d:'))}、节点 {sum(1 for x in sc if x['key'].startswith('x:'))}；"
          f"不评分 {len(rows) - len(sc)}")
    for label, part in (("评分范围·藏品", [x for x in sc if x["level"] == "藏品"]), ("节点", nodes)):
        L = [len(x["desc"]) for x in part] or [0]
        r.say(f"  {label}：简介中位 {statistics.median(L)} 字、最短 {min(L)}；有图 {sum(1 for x in part if x['image'])}/{len(part)}；"
              f"有尺寸 {sum(1 for x in part if x['size'])}；有年代 {sum(1 for x in part if x['year'] and x['year'] not in ('其他', '年代不详'))}")
    r.say(f"  评分范围·藏品的类目 {Counter(x['cat'] or '（空）' for x in sc if x['level'] == '藏品').most_common()}")
    r.say(f"陈列状态 {dict(Counter(x['on'] for x in rows))}")
    r.say(f"展厅 {dict(Counter(x['hall'] or '（空）' for x in rows))}")
    for x in nodes:
        r.say(f"  节点 {x['seq']:>5} [{x['level']}] {x['name']} | {x['hall']} | {x['on']} | 简介 {len(x['desc'])} 字 | {x['twin_why']}")
    for tid, why in TWIN_REJECTED.items():
        r.say(f"  未合并（证据不足）：镇馆之宝 {tid} {why}")
    al = [x for x in objs if x["alias"] and x["key"].startswith("d:")]
    r.say(f"  列表名与详情名不同、详情名进了「别名」的 {len(al)} 件（吴大澂作品：列表是具体题名，详情是名录通称）")
    odd = [x["name"] for x in al if ("篆书" in x["name"]) != ("篆书" in x["alias"]) or
           re.search(r"[七八九]言", x["name"]) and re.search(r"[七八九]言", x["alias"]) and
           re.search(r"[七八九]言", x["name"]).group() != re.search(r"[七八九]言", x["alias"]).group()]
    if odd:
        r.say(f"  ⚠ 两个名字自相矛盾（书体或字数不同），照录，留给审计：{odd}")
    ws = [x["name"] for x in rows if x["name"] != x["name"].strip() or not x["name"]]
    if ws:
        r.say(f"  ⚠ 名称为空或带首尾空白：{ws[:5]}")


# ---------------------------------------------------------------- 写 Excel
HEADER = ["序号", "展厅", "展品名称", "展品简介", "展品图片", "陈列状态", "官方页面", "Tier",
          "对象层级", "类目", "英文名称", "英文简介", "官网ID", "出处",
          "评分范围", "年代", "质地", "尺寸", "文物级别", "别名", "可移动文物编号", "Wikidata", "备注"]


def write_xlsx(rows: list[dict], r) -> None:
    import openpyxl
    from openpyxl.styles import Alignment, Font
    wb = openpyxl.Workbook(write_only=False)
    ws = wb.active
    ws.title = "展品清单"
    objs = sum(1 for x in rows if x["level"] == "藏品")
    sc = sum(1 for x in rows if x["scope"] == SCOPE_YES)
    ws.append(["吉林省博物院 展品清单"])
    ws.append([f"资料整理日期：{dt.date.today():%Y 年 %m 月} | 数据来源：{BASE} 官方接口 | "
               f"藏品 {objs} 件、节点 {len(rows) - objs} 个、评分范围 {sc} 行 | 由 jlpm_build.py 生成，勿手工编辑"])
    ws.append([])
    ws.append(HEADER)
    for c in ws[4]:
        c.font = Font(bold=True)
    for x in sorted(rows, key=lambda x: x["seq"]):
        ws.append([x["seq"], x["hall"], x["name"], x["desc"], x["image"], x["on"], x["url"], "",
                   x["level"], x["cat"], "", "", x["oid"], x["src"],
                   x["scope"], x["year"], x["texture"], x["size"], x["grade"], x["alias"],
                   x["relic"], x["qid"], x["twin_why"]])
    for col, width in zip("ABCDEFGHIJKLMNOPQRSTUVW",
                          (7, 14, 30, 60, 30, 22, 30, 6, 10, 14, 8, 8, 26, 22, 12, 10, 10, 20, 8, 24, 24, 12, 40)):
        ws.column_dimensions[col].width = width
    for row in ws.iter_rows(min_row=5, max_row=4 + sc + 20):
        for c in row:
            c.alignment = Alignment(wrap_text=True, vertical="top")

    info = wb.create_sheet("参观实用信息")
    info.append(["吉林省博物院 — 参观实用信息"])
    info.append([])
    info.append(["项目", "内容"])
    for k, v in (
        ("馆址", "主馆：长春市净月国家高新技术产业开发区永顺路 1666 号（轨道交通 6 号线省博物院站）。"
                 "另一处：「吉林省近现代史展」在长春市北京大街西历史文化街区 A6、A7 栋（春京西），2025-08-19 开展"),
        ("数据范围", "官网「镇馆之宝」17 件；「藏品数据库」17709 件全部入库，只有写了介绍的进评分；"
                     "当前在展的实体展览（馆方标记展出中），以及「白山松水的记忆」基本陈列"),
        ("数据来源", f"{BASE} 的 /api 接口（2026-09-26 至 27 抓取，jlpm_site_scrape.py）。境外网络连不上，须在国内网络抓。"
                     "旧域名 jlmuseum.org 已被占用，不是博物院"),
        ("陈列状态", "官网给不出单件藏品在哪个展厅：展品清单接口全部为空，镇馆之宝的「曾经展出」对每件都是同一串"
                     "最新展览。藏品一律「在展状态未知」。展览的状态照录馆方标记"),
        ("白山松水的记忆", "馆方状态标记为「已结束」，时间栏却写「正在展出」，两处矛盾，记为「未核实」"),
        ("身份核对", "名称与国家文物局《全国馆藏文物名录》（经 Wikidata）逐字相同且两边唯一的，挂可移动文物编号"),
        ("图片版权", "图片地址指向吉林省博物院与吉林省数字博物馆的服务器，版权归原单位，仅供个人查阅"),
    ):
        info.append([k, v])
    info.column_dimensions["A"].width = 14
    info.column_dimensions["B"].width = 100
    XLSX.parent.mkdir(exist_ok=True)
    wb.save(XLSX)
    r.say(f"\n写出 {XLSX.relative_to(HERE)}（{len(rows)} 行）")


class Report:
    def __init__(self):
        self.lines: list[str] = []

    def say(self, *a) -> None:
        line = " ".join(str(x) for x in a)
        self.lines.append(line)
        print(line, flush=True)


def main() -> None:
    sys.stdout.reconfigure(encoding="utf-8")        # 中文 Windows 控制台默认 GBK；
    sys.stderr.reconfigure(encoding="utf-8")        # sys.exit 的报错走 stderr，也要改
    ap = argparse.ArgumentParser(description=__doc__.split("\n\n")[0])
    ap.add_argument("--check", action="store_true", help="只做合并、名录核对与体检，不写文件")
    args = ap.parse_args()
    r = Report()
    r.say(f"# 吉林省博物院 源表生成{'（只检查）' if args.check else ''}")
    rows = build_rows(r)
    assign_seq(rows, write=not args.check, r=r)
    match_catalog(rows, r)
    health(rows, r)
    if not args.check:
        write_xlsx(rows, r)
        REPORT.write_text("\n".join(r.lines) + "\n", "utf-8")


if __name__ == "__main__":
    main()
