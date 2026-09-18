#!/usr/bin/env python3
"""开 Google 搜索工具，逐件查具体展厅号，**只写有出处的**。

用法：
    python3 gallery_grounded.py --dry-run          # 只看选中哪些行、不调 API
    python3 gallery_grounded.py --limit 2          # 先跑 2 件
    python3 gallery_grounded.py                    # S 级在展且没查到展厅的全部

**为什么有这个脚本（2026-09-15）**

`fill_onview_gemini.py` 走的是裸 `generateContent`，**没开任何工具**，模型只能凭记忆答。
实测同一个问题（Homer《The Fog Warning》在哪个展厅）：

  · 不开搜索 → `Gallery 222`，并自称「根据波士顿美术馆的官方馆藏记录」（它没查过任何记录）
  · 开搜索（用户在 Gemini App 里）→ `Gallery 234`

**两个数字不一样，而不开搜索的那个还伪造了出处。** 所以那个脚本的提示词写死「记不准
展厅号就留空」——挡住的正是 `Gallery 222` 这类东西，代价是展厅列大片为空。

本脚本换一条路：**给它检索能力，然后要求出处**。规则只有一条硬的 ——

    **没有 groundingChunks（检索来源）就不写。** 模型凭记忆答出来的一律丢弃。

这跟仓库里反复踩的「取不到就退而求其次」是反着来的：取不到就留空并喊出来。

**结构化输出与搜索工具不能同时用**（responseSchema 会和 tools 冲突），故改为让模型
以 `GALLERY: xxx` 结尾，正则取值；取不到就算失败，不猜。

**证据落仓库**：每次调用的答案原文、解析结果、来源 URL、实际搜索词全部写进
`gallery_grounded.jsonl` —— 抓回来的东西要落进仓库，不能只落进 Excel。
"""
from __future__ import annotations

import argparse
import json
import pathlib
import re
import shutil
import sys
import time
import urllib.error
import urllib.request

import openpyxl

from gallery_text import SPECIFIC_RE, is_specific
from tls import ssl_ctx
from merged_xlsx import merged_xlsx

HERE = pathlib.Path(__file__).parent
OUT = HERE / "gallery_grounded.jsonl"
SHEET = "去重后总表"

MODEL = "gemini-3.5-flash"
ENDPOINT = ("https://generativelanguage.googleapis.com/v1beta/models/"
            "{model}:generateContent")
KEY_FILE = "~/.gemini_key"
MIN_GAP = 8.0          # 秒；带搜索的调用更贵，放慢一点

GALLERY_RE = re.compile(r"GALLERY:\s*(.+?)\s*$", re.I | re.M)
ONVIEW_RE = re.compile(r"ONVIEW:\s*(YES|NO|UNKNOWN)\s*$", re.I | re.M)

PROMPT = """请查证波士顿美术馆（Museum of Fine Arts, Boston）这件藏品的当前状态。

藏品：{name}
{acc}

依次回答两件事：
1. 它现在是否在 MFA 的展厅公开展出？
2. 若在展，具体在哪个展厅？

要求：
1. **必须用搜索查证**，优先采信 mfa.org / collections.mfa.org 的页面。
   collections.mfa.org 的单件页上有 "On View" 或 "Not on View" 字样，那是最可靠的依据。
2. 只有查到**具体展厅**（如 "Gallery 234"）才算数。
   「美洲艺术翼」「古埃及展区」这类是**区域不是展厅**，不算。
3. **查不到就如实说查不到，不要凭印象回答。**
   这件东西可能早已不属于 MFA（易主、退藏、长期外借），也可能在库房轮换。
   凭「这类作品通常陈列在某厅」推断具体某件在不在，正是要避免的错误。
4. 判断在展与否，说的是**这一件**，不是同一位作者的其他作品。
5. **检索次数就是成本**：请用尽量少、尽量精准的关键词查证，
   能一次查清就不要拆成多次搜索。

最后另起两行，严格用这两个格式收尾：
ONVIEW: YES / NO / UNKNOWN
GALLERY: Gallery 234 / NONE"""


def source_label(sources: list[dict]) -> str:
    """给来源起个短标签写进单元格。

    ⚠ Gemini 的 groundingChunks.uri 多半是重定向地址
    （vertexaisearch.cloud.google.com/grounding-api-redirect/…），解析主机名得到的
    是 Google 的域名而不是真正的出处。真域名通常在 `title` 里，故优先取 title。
    官方源排在前面，让单元格里一眼看得出可信度。
    """
    titles = [str(s.get("title") or "").strip() for s in sources]
    titles = [t for t in titles if t]
    for t in titles:                       # 官方源优先
        if "mfa.org" in t.lower():
            return t
    if titles:
        return titles[0]
    m = re.search(r"//([^/]+)", sources[0]["uri"])
    return re.sub(r"^www\.", "", m.group(1)) if m else "未知来源"



def key(key_file: str = KEY_FILE) -> str:
    """读 API key。

    ⚠ 本机同时存着两个 key：`~/.gemini_key`（付费项目）与 `~/.gemini_key_free`。
    带检索的调用要走付费项目 —— 免费层没有 Search grounding 的额度。
    **哪个 key 花钱是要写清楚的事**，所以留 --key-file 显式指定，不靠默认值猜。
    """
    p = pathlib.Path(key_file).expanduser()
    if not p.exists():
        sys.exit(f"找不到 {p}")
    return p.read_text().strip()


def ask(name: str, acc, model: str, api_key: str) -> dict:
    """一次带搜索的调用。返回 {text, gallery, sources, queries}。"""
    body = {
        "contents": [{"role": "user", "parts": [{"text": PROMPT.format(
            name=name,
            acc=f"馆藏号：{acc}" if acc else "（无馆藏号）")}]}],
        "tools": [{"google_search": {}}],
    }
    req = urllib.request.Request(
        ENDPOINT.format(model=model), data=json.dumps(body).encode(),
        headers={"x-goog-api-key": api_key, "Content-Type": "application/json"})
    with urllib.request.urlopen(req, timeout=300, context=ssl_ctx()) as r:
        resp = json.load(r)

    cand = resp["candidates"][0]
    text = "".join(p.get("text", "") for p in cand.get("content", {}).get("parts", []))
    gm = cand.get("groundingMetadata") or {}
    sources = []
    for c in (gm.get("groundingChunks") or []):
        w = c.get("web") or {}
        if w.get("uri"):
            sources.append({"title": w.get("title"), "uri": w["uri"]})
    m = GALLERY_RE.search(text)
    raw = m.group(1).strip() if m else ""
    sm = SPECIFIC_RE.search(raw)
    om = ONVIEW_RE.search(text)
    return {
        "text": text,
        "onview_raw": om.group(1).upper() if om else "",
        "gallery_raw": raw,
        # 收尾格式取不到就算失败，不去正文里猜 —— 猜出来的东西没法复核。
        "gallery": sm.group(0) if sm else None,
        "sources": sources,
        "queries": gm.get("webSearchQueries") or [],
        "usage": resp.get("usageMetadata"),
    }


def main() -> None:
    ap = argparse.ArgumentParser()
    ap.add_argument("--tier", default="S")
    ap.add_argument("--limit", type=int)
    ap.add_argument("--model", default=MODEL)
    ap.add_argument("--dry-run", action="store_true")
    ap.add_argument("--xlsx", help="要改的工作簿；不给则按 merged_xlsx() 解析合并版。"
                                   "**就地改**，所以跑新一轮时建议先 copy_sheet.py 抄一份副本。")
    ap.add_argument("--max-queries", type=int,
                    help="搜索次数预算。Grounding 按**模型实际发起的搜索次数**计费"
                         "（groundingMetadata.webSearchQueries 的长度），不是按请求数 —— "
                         "一次请求可能触发多次搜索。累计达到这个数就停下并保存已有结果。")
    ap.add_argument("--key-file", default=KEY_FILE,
                    help=f"API key 文件，默认 {KEY_FILE}（付费项目）。"
                         "带检索的调用只有付费项目跑得了。")
    args = ap.parse_args()

    XLSX = pathlib.Path(args.xlsx) if args.xlsx else merged_xlsx()
    if args.xlsx and not XLSX.exists():
        sys.exit(f"--xlsx 指定的文件不存在：{XLSX}")

    wb = openpyxl.load_workbook(XLSX)
    ws = wb[SHEET]
    hdr = [c.value for c in ws[1]]
    for n in ("Tier", "陈列状态", "展厅", "展品名称", "馆藏号"):
        if n not in hdr:
            sys.exit(f"表头缺列「{n}」")
    c_t, c_s = hdr.index("Tier") + 1, hdr.index("陈列状态") + 1
    c_g, c_n = hdr.index("展厅") + 1, hdr.index("展品名称") + 1
    c_a = hdr.index("馆藏号") + 1

    c_seq = hdr.index("序号") + 1 if "序号" in hdr else None
    c_mus = hdr.index("博物馆") + 1 if "博物馆" in hdr else None

    # 已核实离馆的展品（Wikidata P195+P582，见 verify_mfa_membership.py）。
    # **核实过的事实压过模型推断** —— 模型曾把其中 2 件判成「在展」，
    # 莫奈《The Fort of Antibes》还配了 Gallery 252（那个号本身没错，
    # MFA 的莫奈确实在 252 厅，错的是这一件已于 2011 年易主 Museum Barberini）。
    # 所以这些行一律不查也不动，免得一次调用就把核实结果冲掉。
    departed = set()
    mp = HERE / "mfa_membership.json"
    if mp.exists():
        departed = {r["seq"] for r in json.loads(mp.read_text(encoding="utf-8"))["left_mfa"]}

    tiers = {x.strip() for x in args.tier.split(",") if x.strip()}
    todo, protected = [], 0
    for row in range(2, ws.max_row + 1):
        if ws.cell(row, c_t).value not in tiers:
            continue
        g = str(ws.cell(row, c_g).value or "")
        if is_specific(g):                    # 已有能走过去的定位，不动
            continue
        if c_seq and c_mus and "扩充清单" in str(ws.cell(row, c_mus).value or ""):
            if ws.cell(row, c_seq).value in departed:
                protected += 1
                continue
        todo.append({"row": row, "name": str(ws.cell(row, c_n).value),
                     "acc": ws.cell(row, c_a).value, "cur": g,
                     "state": ws.cell(row, c_s).value})
    if args.limit:
        todo = todo[:args.limit]

    print(f"待查 {len(todo)} 行（Tier {args.tier}，展厅给不出具体位置）"
          f"｜已核实离馆而跳过 {protected} 行")
    for t in todo[:40]:
        print(f"  行{t['row']:<6} {t['name'][:42]:<44} 状态={str(t['state'] or '空'):<5} "
              f"展厅={t['cur'][:16] or '空'}")
    if len(todo) > 40:
        print(f"  …… 其余 {len(todo)-40} 行略")
    if not args.dry_run and not args.max_queries and len(todo) > 200:
        # 大批量又不设预算，是这个仓库里最容易出事的组合：按搜索次数计费，
        # 而搜索次数事前不可知（一次请求可能拆成多次搜索）。宁可先喊一声。
        print(f"\n⚠ 待查 {len(todo)} 行且未设 --max-queries。"
              f"Grounding 按搜索次数计费，一次请求可能触发多次搜索，"
              f"实际次数会高于 {len(todo)}。\n"
              f"  建议先小批验证，或用 --max-queries 设上限。")
        sys.exit("已中止：请显式给出 --max-queries（或用 --limit 先小批跑）。")

    if args.dry_run or not todo:
        return

    api_key = key(args.key_file)
    wrote_gal, wrote_ov, nosrc, nogal, failed = [], [], [], [], []
    n_queries = 0          # 累计搜索次数 —— 这才是计费口径
    stopped = None
    last = 0.0
    for t in todo:
        if args.max_queries and n_queries >= args.max_queries:
            stopped = f"已达搜索次数预算 {args.max_queries}（实际 {n_queries}）"
            print(f"\n⏹ {stopped}，停止；已完成的结果照常保存。")
            break
        gap = MIN_GAP - (time.time() - last)
        if gap > 0:
            time.sleep(gap)
        last = time.time()
        try:
            r = ask(t["name"], t["acc"], args.model, api_key)
        except urllib.error.HTTPError as e:
            msg = e.read().decode("utf-8", "replace")[:160]
            print(f"  行{t['row']} 调用失败 HTTP {e.code} {msg}")
            failed.append(t["row"])
            continue
        except Exception as e:
            print(f"  行{t['row']} 调用失败 {type(e).__name__} {str(e)[:120]}")
            failed.append(t["row"])
            continue

        nq = len(r["queries"])
        n_queries += nq
        rec = {"row": t["row"], "name": t["name"], "acc": str(t["acc"]),
               "model": args.model, "n_queries": nq, "n_queries_total": n_queries,
               "before": {"state": t["state"], "gallery": t["cur"]},
               **r}
        with OUT.open("a", encoding="utf-8") as f:
            f.write(json.dumps(rec, ensure_ascii=False) + "\n")

        # 硬规则：没有检索来源，什么都不写 —— 在展与否也不写。
        # 凭记忆答的在展状态正是上一轮被证伪的那批。
        if not r["sources"]:
            print(f"  行{t['row']} ⚠ 无检索来源，整条丢弃"
                  f"（它说 ONVIEW={r['onview_raw'] or '?'} GALLERY={r['gallery_raw']!r}）")
            nosrc.append(t["row"])
            continue

        src = source_label(r["sources"])
        ov = r["onview_raw"]
        if ov == "YES":
            ws.cell(t["row"], c_s).value = "在展"
            wrote_ov.append((t["row"], "在展"))
            if r["gallery"]:
                ws.cell(t["row"], c_g).value = f"{r['gallery']}（来源：{src}）"
                wrote_gal.append((t["row"], r["gallery"], len(r["sources"]), t["name"][:34]))
                print(f"  行{t['row']} ✓ 在展 {r['gallery']}  来源 {len(r['sources'])} 条")
            else:
                nogal.append(t["row"])
                print(f"  行{t['row']} · 在展，但未查到具体展厅（原文 {r['gallery_raw']!r}）")
        elif ov == "NO":
            ws.cell(t["row"], c_s).value = "不在展"
            ws.cell(t["row"], c_g).value = None
            wrote_ov.append((t["row"], "不在展"))
            print(f"  行{t['row']} ✓ 不在展（来源 {len(r['sources'])} 条）")
        else:
            # UNKNOWN 或没按格式收尾：**不覆盖**已有状态。
            # 把「查不到」写成「未知」会把一条已有判断降级成无信息。
            nogal.append(t["row"])
            print(f"  行{t['row']} · 查不到（ONVIEW={ov or '未按格式收尾'}），保持原值")

    done = len(wrote_gal) + len(wrote_ov) + len(nogal) + len(nosrc)
    print(f"\n写入展厅 {len(wrote_gal)}｜写入在展状态 {len(wrote_ov)}｜"
          f"查不到 {len(nogal)}｜无出处丢弃 {len(nosrc)}｜调用失败 {len(failed)}")
    print(f"搜索次数合计 {n_queries} 次"
          + (f"，平均每件 {n_queries/done:.2f} 次" if done else "")
          + "（计费口径；请求数 != 搜索数）")
    if stopped:
        print(f"⏹ 本轮提前结束：{stopped}")
    for row, g, n, nm in wrote_gal:
        print(f"  行{row:<6} {g:<14} 来源{n}条  {nm}")
    if len(wrote_gal) + len(wrote_ov) + len(nogal) + len(nosrc):
        print(f"\n逐条证据已写入 {OUT.name}")

    if not (wrote_gal or wrote_ov):
        print("没有任何行被写入，不重存文件。")
        return
    bak = XLSX.with_suffix(".xlsx.bak4")
    if not bak.exists():
        shutil.copy2(XLSX, bak)
        print(f"改动前已备份到 {bak.name}")
    wb.save(XLSX)
    print(f"已写回 {XLSX.name}")


if __name__ == "__main__":
    main()
