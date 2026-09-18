#!/usr/bin/env python3
"""把部门/翼级的展厅说法标注清楚，并从模型 note 里救回被丢掉的那一类。零 API。

**为什么要这一步（2026-09-14，用户定）**

`fill_onview_gemini.py` 的提示词写了「记不准展厅号就把 gallery 留空」，于是模型
把「古埃及展厅」「印象派展厅」这类说法放进了 `note`，`gallery` 返回空串 ——
表里「展厅」列就是空的（用户问「既然查到是古埃及展厅，为什么展厅列为空」即此）。

这类说法的问题不是没用，是**精度与真定位不同级**：MFA 的古埃及展区横跨多个厅，
「古埃及展厅」走不到具体位置。若直接写进「展厅」列，它会和 `Gallery 252` 长得一样，
事后分不出精度 —— 与表里那 78 行「Art of Europe 欧洲艺术」同一个形状，
也正是北斋《赤富士》被误判在展的来源。

故：**保留信息，但把精度写在值里**，加后缀 `（部门级，未确认具体展厅）`。

两件事：
  1. 在展 + 展厅为空 + note 里有展厅说法 → 写入带后缀的值；
  2. 模型已写进展厅、但取值是翼/部门级的 → 就地补后缀。

**什么算「具体位置」**（这几类不加后缀，其余模型给的值一律算部门级）：
  · `Gallery <数字>`；
  · MFA 实际的命名空间：Sargent Colonnade / Sargent Rotunda / Huntington Avenue Plaza。
认不出的按部门级处理 —— 宁可多标一句「未确认」，不可让粗粒度冒充精确定位。

只动本轮模型写过的行（以 `onview_cache.jsonl` 为准），不碰源数据里原有的展厅值。
"""
from __future__ import annotations

import json
import pathlib
import re
import shutil
import sys

import openpyxl

from merged_xlsx import merged_xlsx

HERE = pathlib.Path(__file__).parent
CACHE = HERE / "onview_cache.jsonl"
SHEET = "去重后总表"

SUFFIX = "（部门级，未确认具体展厅）"

# 具体位置的白名单：匹配上就不加后缀
SPECIFIC = re.compile(
    r"Gallery\s*\d+"
    r"|Sargent\s+(?:Colonnade|Rotunda)"
    r"|Huntington\s+Avenue\s+Plaza",
    re.I,
)

# 从 note 里提取展厅说法。中英各一套：
#   中文——一段不含标点的、以「展厅/展馆」收尾的短语；
#   英文——note 常写成叙述句（"permanently on view in the Art of the Americas Wing"），
#          只认 "Art of ... Wing" 这种馆方正式称谓，认不出就不写。
NOTE_GALLERY = re.compile(r"([^，。；、\s][^，。；、]{1,17}(?:展厅|展馆))")
NOTE_GALLERY_EN = re.compile(r"(Art of the [A-Za-z]+ Wing|Art of (?:Europe|Asia|the Americas))", re.I)


def strip_prefix(phrase: str) -> str:
    """去掉开头的馆名前缀。

    ⚠ 不要用 `lstrip("MFA")` —— str.lstrip 按**字符集**剥，
    "Art of the Americas Wing" 开头的 A 在集合里，会被吃成 "rt of…"。
    """
    out = phrase.strip()
    for pre in ("MFA ", "MFA"):
        if out.startswith(pre):
            out = out[len(pre):]
            break
    return out.strip()


def derive(a: dict) -> str | None:
    """从模型 note 推出该写的展厅说法；推不出返回 None。"""
    note = a.get("note") or ""
    m = NOTE_GALLERY.search(note) or NOTE_GALLERY_EN.search(note)
    return strip_prefix(m.group(1)) if m else None


def main() -> None:
    XLSX = merged_xlsx()
    if not CACHE.exists():
        sys.exit(f"找不到 {CACHE}")
    ans = {}
    for line in CACHE.read_text(encoding="utf-8").splitlines():
        if line.strip():
            for it in json.loads(line)["resp"]["items"]:
                if "id" in it:
                    ans[int(it["id"])] = it
    print(f"缓存里模型答过 {len(ans)} 行")

    wb = openpyxl.load_workbook(XLSX)
    ws = wb[SHEET]
    hdr = [c.value for c in ws[1]]
    for n in ("陈列状态", "展厅", "展品名称"):
        if n not in hdr:
            sys.exit(f"表头缺列「{n}」")
    c_st = hdr.index("陈列状态") + 1
    c_gal = hdr.index("展厅") + 1
    c_name = hdr.index("展品名称") + 1

    recovered, marked, fixed = [], [], []
    for row, a in ans.items():
        if row > ws.max_row:
            continue
        if ws.cell(row, c_st).value != "在展":     # 已被离馆修正等改掉的不碰
            continue
        cur = ws.cell(row, c_gal).value
        name = str(ws.cell(row, c_name).value)[:40]

        if cur:
            # 情形 2：模型写过值，判断精度
            if str(cur).endswith(SUFFIX):
                # 之前写过后缀的，复核一遍本体对不对（曾因 lstrip 把首字母吃掉）
                want = derive(a)
                if want and str(cur) != f"{want}{SUFFIX}":
                    ws.cell(row, c_gal).value = f"{want}{SUFFIX}"
                    fixed.append((row, str(cur), f"{want}{SUFFIX}", name))
                continue
            if SPECIFIC.search(str(cur)):
                continue
            if (a.get("gallery") or "").strip() != str(cur).strip():
                continue                            # 不是模型写的，别动
            ws.cell(row, c_gal).value = f"{cur}{SUFFIX}"
            marked.append((row, str(cur), name))
        else:
            # 情形 1：展厅为空，从 note 里救
            phrase = derive(a)
            if not phrase:
                continue
            ws.cell(row, c_gal).value = f"{phrase}{SUFFIX}"
            recovered.append((row, phrase, name))

    print(f"\n从 note 救回 {len(recovered)} 行：")
    for row, p, nm in recovered:
        print(f"  行{row:<6} {p:<22} {nm}")
    print(f"\n给翼/部门级取值补后缀 {len(marked)} 行：")
    for row, p, nm in marked:
        print(f"  行{row:<6} {p:<34} {nm}")

    print(f"\n纠正写坏的取值 {len(fixed)} 行：")
    for row, old, new, nm in fixed:
        print(f"  行{row:<6} {old!r} → {new!r}  {nm}")

    if not (recovered or marked or fixed):
        print("\n没有需要改的行，不重存文件。")
        return

    bak = XLSX.with_suffix(".xlsx.bak3")
    if not bak.exists():
        shutil.copy2(XLSX, bak)
        print(f"\n改动前已备份到 {bak.name}")
    wb.save(XLSX)
    print(f"已写回 {XLSX.name}")


if __name__ == "__main__":
    main()
