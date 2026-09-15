#!/usr/bin/env python3
"""按展品名称向 Gemini 查在展状态与展厅，就地写回合并版 Excel。

用法：
    python3 fill_onview_gemini.py --dry-run          # 不调 API，只验通路与回写
    python3 fill_onview_gemini.py --limit 20         # 先跑 20 件看看
    python3 fill_onview_gemini.py                    # 全量

**⚠ 这份数据的性质必须写在这里，因为表里看不出来。**

写进 `陈列状态` / `展厅` 两列的值来自**语言模型的知识**，不是 MFA 的馆藏系统。
模型没有实时在展数据，展厅号与在展状态是逐周变动的。已知的三条硬事实：

  · `collections.mfa.org` 有人机验证 CAPTCHA，robots 禁 `/search`、`Crawl-delay: 30`
    —— 权威源无法程序化访问（2026-09-14 实测，无头浏览器同样被挡）；
  · MFA 没有公开 API（哈佛、大都会有）；
  · Wikidata 的 P276 只到馆一级，没有展厅、没有在展状态。

所以这两列**填上之后，与那 56 行有真实依据的行在表里长得一模一样**，事后分不出来。
表里原有的 312 行「在展」中：56 行写明了具体展厅或「常设展出」（有依据）、
78 行只写了部门/专辑（无依据）、178 行照抄源文件且时效不明。
本脚本默认**只碰「未知」与空行**（`--overwrite` 才动已有值），就是为了不把那 56 行盖掉。

用户 2026-09-14 在知情后决定采用这条路。脚本照办，但把上面这段留在这里。

**缓存**：结果逐条落 `onview_cache.jsonl`，键含型号与提示词全文 —— 与 `llm_cache.py`
同一个约定：提示词一字未改必然命中（不重复付费），改了必然不命中（拿不到旧判据的答案）。
本脚本不走 MySQL（这台机器没有 `~/.my.cnf`），故用本地 JSONL。
"""
from __future__ import annotations

import argparse
import hashlib
import json
import pathlib
import shutil
import sys
import time

import openpyxl

import gemini_api

XLSX = pathlib.Path(__file__).parent / "export" / "展品_波士顿美术馆_合并.xlsx"
SHEET = "去重后总表"
CACHE = pathlib.Path(__file__).parent / "onview_cache.jsonl"

MODEL = "gemini-3.6-flash"     # 2.5-flash 对新用户已下线，官方指向这个

SYSTEM = """你在为一份波士顿美术馆（MFA Boston）的展品清单核对陈列状态。

对每一件展品，回答两件事：
1. on_view —— 它目前是否在 MFA 的常设展厅公开展出。
2. gallery —— 若在展，具体展厅（例如 "Gallery 252" 或 "Art of the Americas Wing, Level 2"）。

**判据（严格遵守）**

· 你没有 MFA 的实时馆藏数据。**只有在你确知这件作品是该馆长期陈列的知名展品时**，
  才回答 on_view=true；否则一律 "unknown"。
· **不要猜展厅号。** 记不准具体展厅就把 gallery 留空，只要 on_view 判断。
  写错的展厅号比空着有害得多 —— 游客会照着它走过去然后扑空。
· 库房藏品、纸本（浮世绘、素描、摄影）等因保存需要轮换展出的门类，
  除非你确知它常年陈列，否则一律 "unknown"。
· 已知离馆、易主、或有捐赠条款限制不得公开展出的，回答 on_view=false 并在 note 说明。
· confidence 如实填：high 只用于你确信的知名常设展品。

宁可答 unknown，不要编。"""

SCHEMA = {
    "type": "object",
    "properties": {
        "items": {
            "type": "array",
            "items": {
                "type": "object",
                "properties": {
                    # ⚠ 标识用 xlsx 行号，不用「序号」——
                    # 这张表合并了两个 museum key，两边各自从 1 开始编号，
                    # `序号` 在表内不唯一（AGENTS.md 的软键是 (museum_key, source_seq)）。
                    # 拿 seq 建 行映射会让 A 件的答案写进 B 件的行，且不报错。
                    "id": {"type": "integer"},
                    "on_view": {"type": "string", "enum": ["true", "false", "unknown"]},
                    "gallery": {"type": "string"},
                    "confidence": {"type": "string", "enum": ["high", "medium", "low"]},
                    "note": {"type": "string"},
                },
                "required": ["id", "on_view", "gallery", "confidence"],
            },
        }
    },
    "required": ["items"],
}


def load_cache() -> dict:
    if not CACHE.exists():
        return {}
    out = {}
    for line in CACHE.read_text(encoding="utf-8").splitlines():
        if line.strip():
            r = json.loads(line)
            out[r["key"]] = r["resp"]
    return out


def cache_key(model: str, system: str, user: str) -> str:
    h = hashlib.sha256()
    for part in (model, system, user, json.dumps(SCHEMA, sort_keys=True)):
        h.update(part.encode())
        h.update(b"\x00")
    return h.hexdigest()


def build_prompt(batch: list[dict]) -> str:
    lines = ["以下是展品清单，逐件作答，id 必须原样返回：", ""]
    for it in batch:
        acc = f"，馆藏号 {it['acc']}" if it["acc"] else ""
        lines.append(f"id={it['row']}｜{it['name']}{acc}")
    return "\n".join(lines)


def main() -> None:
    ap = argparse.ArgumentParser()
    ap.add_argument("--batch", type=int, default=20, help="每次调用几件")
    ap.add_argument("--limit", type=int, help="只处理前 N 件")
    ap.add_argument("--model", default=MODEL)
    ap.add_argument("--tier", default="",
                    help="只跑这些 Tier，逗号分隔，如 S,A（默认全部）")
    ap.add_argument("--overwrite", action="store_true",
                    help="连已有依据的行也覆盖（默认只填「未知」与空行）")
    ap.add_argument("--dry-run", action="store_true",
                    help="不调 API，验证读表/回写通路")
    args = ap.parse_args()

    if not XLSX.exists():
        sys.exit(f"找不到 {XLSX}")

    wb = openpyxl.load_workbook(XLSX)
    if SHEET not in wb.sheetnames:
        sys.exit(f"没有 sheet「{SHEET}」，现有：{wb.sheetnames}")
    ws = wb[SHEET]

    hdr = [c.value for c in ws[1]]
    need = ("序号", "展品名称", "馆藏号", "陈列状态", "展厅", "Tier")
    col = {}
    for name in need:
        if name not in hdr:
            sys.exit(f"表头里找不到列「{name}」，现有前 10 列：{hdr[:10]}")
        col[name] = hdr.index(name) + 1        # openpyxl 是 1-based

    tiers = {t.strip() for t in args.tier.split(",") if t.strip()}
    todo = []
    for row in range(2, ws.max_row + 1):
        if tiers and ws.cell(row, col["Tier"]).value not in tiers:
            continue
        status = ws.cell(row, col["陈列状态"]).value
        if not args.overwrite and status not in (None, "", "未知"):
            continue
        name = ws.cell(row, col["展品名称"]).value
        if not name:
            continue
        todo.append({
            "row": row,
            "seq": ws.cell(row, col["序号"]).value,
            "name": str(name),
            "acc": ws.cell(row, col["馆藏号"]).value,
        })
    if args.limit:
        todo = todo[:args.limit]
    scope = f"Tier {args.tier}" if tiers else "全部 Tier"
    print(f"待处理 {len(todo)} 行（{scope}；全表 {ws.max_row - 1} 行）")
    if not todo:
        sys.exit("没有需要填的行")

    if args.dry_run:
        print("--dry-run：不调 API。样例 5 行 —")
        for t in todo[:5]:
            print(f"  行{t['row']} seq={t['seq']} {t['name'][:60]}")
        print("\n提示词样例：\n" + build_prompt(todo[:3])[:600])
        return

    if not gemini_api.available():
        sys.exit("gemini_api.available() 为 False —— 检查 ~/.gemini_key 是否存在且权限正确")

    cache = load_cache()
    filled = skipped = 0

    for i in range(0, len(todo), args.batch):
        batch = todo[i:i + args.batch]
        user = build_prompt(batch)
        key = cache_key(args.model, SYSTEM, user)
        if key in cache:
            resp = cache[key]
            print(f"[{i//args.batch + 1}] 命中缓存")
        else:
            try:
                resp, usage = gemini_api.ask(SYSTEM, user, SCHEMA, model=args.model)
            except Exception as e:
                print(f"[{i//args.batch + 1}] 调用失败：{type(e).__name__} {str(e)[:160]}")
                time.sleep(10)
                continue
            with CACHE.open("a", encoding="utf-8") as f:
                f.write(json.dumps({"key": key, "resp": resp}, ensure_ascii=False) + "\n")
            print(f"[{i//args.batch + 1}] {len(batch)} 件，"
                  f"out={getattr(usage, 'completion_tokens', '?')}")

        got = {int(r["id"]): r for r in resp.get("items", []) if "id" in r}
        missing = [b["row"] for b in batch if b["row"] not in got]
        if missing:
            # 「取不到就喊」—— 不静默跳过，也不拿别的值顶上
            print(f"    ⚠ 模型漏返 {len(missing)} 件：{missing[:8]}")
        for b in batch:
            r = got.get(b["row"])
            if not r:
                skipped += 1
                continue
            row = b["row"]
            ov = r.get("on_view")
            if ov == "true":
                ws.cell(row, col["陈列状态"]).value = "在展"
                g = (r.get("gallery") or "").strip()
                if g:
                    ws.cell(row, col["展厅"]).value = g
            elif ov == "false":
                ws.cell(row, col["陈列状态"]).value = "不在展"
            else:
                ws.cell(row, col["陈列状态"]).value = "未知"
            filled += 1

    # 一行没填就不要重存 —— openpyxl 重写工作簿会丢图表、条件格式等它不认识的东西，
    # 没有改动却付这个代价是纯亏。（2026-09-14 全批 429 时踩到。）
    if not filled:
        print(f"\n没有任何行被填充（漏返跳过 {skipped} 行），不重存文件。")
        return

    bak = XLSX.with_suffix(".xlsx.bak")
    if not bak.exists():
        shutil.copy2(XLSX, bak)
        print(f"原文件已备份到 {bak.name}")
    wb.save(XLSX)
    print(f"\n已写回 {XLSX.name}：填了 {filled} 行，漏返跳过 {skipped} 行")


if __name__ == "__main__":
    main()
