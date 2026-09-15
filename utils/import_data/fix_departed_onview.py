#!/usr/bin/env python3
"""把已离馆展品的陈列状态钉成「不在展」，并清掉展厅。零 API。

**为什么需要这一步（2026-09-14）**

`fill_onview_gemini.py` 让模型按名称判断在展状态。在我们手上有独立答案的样本上，
它判错了，而且错在最有害的方向：

  · `The Fort of Antibes`（莫奈《安提布堡》，ext seq 4118）→ 判「在展」，还给了 `Gallery 252`
  · `The Annunciation`（埃尔·格列柯，ext seq 138）→ 判「在展」

这两件都在 `mfa_membership.json` 的 `left_mfa` 里 —— 2026-09-11 用 Wikidata
P195（收藏机构）+ P582（离馆时间）核实过，**已经不在 MFA**。
模型给一件已经易主的画配上具体展厅号，游客会照着走过去然后扑空。

**核实过的事实优先于模型的推断。** 本脚本用 `left_mfa` 的 32 条覆盖那两列。
匹配键是 `(博物馆=扩充清单, 序号)` —— `序号` 在合并表里不唯一（两个 museum key
各自从 1 编号），单用 seq 会张冠李戴。
"""
from __future__ import annotations

import json
import pathlib
import shutil
import sys

import openpyxl

HERE = pathlib.Path(__file__).parent
XLSX = HERE / "export" / "展品_波士顿美术馆_合并.xlsx"
MEMBERSHIP = HERE / "mfa_membership.json"
SHEET = "去重后总表"
EXT_MUSEUM = "波士顿美术馆（扩充清单）"


def main() -> None:
    if not MEMBERSHIP.exists():
        sys.exit(f"找不到 {MEMBERSHIP}")
    data = json.loads(MEMBERSHIP.read_text(encoding="utf-8"))
    left = {x["seq"]: x["name"] for x in data["left_mfa"] if x.get("seq") is not None}
    print(f"已核实离馆: {len(left)} 件")

    wb = openpyxl.load_workbook(XLSX)
    ws = wb[SHEET]
    hdr = [c.value for c in ws[1]]
    for name in ("序号", "博物馆", "陈列状态", "展厅", "展品名称"):
        if name not in hdr:
            sys.exit(f"表头缺列「{name}」")
    c_seq = hdr.index("序号") + 1
    c_mus = hdr.index("博物馆") + 1
    c_st = hdr.index("陈列状态") + 1
    c_gal = hdr.index("展厅") + 1
    c_name = hdr.index("展品名称") + 1

    changed = []
    hit = set()
    for row in range(2, ws.max_row + 1):
        if ws.cell(row, c_mus).value != EXT_MUSEUM:
            continue
        seq = ws.cell(row, c_seq).value
        if seq not in left:
            continue
        hit.add(seq)
        before = (ws.cell(row, c_st).value, ws.cell(row, c_gal).value)
        ws.cell(row, c_st).value = "不在展"
        ws.cell(row, c_gal).value = None
        if before != ("不在展", None):
            changed.append((row, str(ws.cell(row, c_name).value)[:46], before))

    missing = set(left) - hit
    if missing:
        # 「取不到就喊」：对不上的要说出来，不能默默少改几行
        print(f"⚠ 有 {len(missing)} 条在表里没找到对应行: {sorted(missing)[:10]}")

    print(f"匹配到 {len(hit)} 件，实际改动 {len(changed)} 行")
    for row, nm, before in changed:
        print(f"  行{row:<6} {before[0]!s:<5} 展厅={str(before[1])[:18]:<20} → 不在展  {nm}")

    if not changed:
        print("没有需要改的行，不重存文件。")
        return

    bak = XLSX.with_suffix(".xlsx.bak2")
    if not bak.exists():
        shutil.copy2(XLSX, bak)
        print(f"\n改动前已备份到 {bak.name}")
    wb.save(XLSX)
    print(f"已写回 {XLSX.name}")


if __name__ == "__main__":
    main()
