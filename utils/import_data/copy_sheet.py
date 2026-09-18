#!/usr/bin/env python3
"""把一个工作簿里的某一个 sheet 原样抄进一个新工作簿。

用法：
    python3 copy_sheet.py --src A.xlsx --sheet 去重后总表 --out B.xlsx

**为什么要有它（2026-09-16）**

`fill_onview_gemini.py` / `gallery_grounded.py` 这几个脚本都是**就地改**工作簿的。
就地改有两个代价：① 一旦写错，原始导出就没了（备份是 `.bak` 系列，容易被下一次
覆盖或搞混）；② openpyxl 重存会丢掉它不认识的东西（图表、条件格式）。
所以要跑新一轮展厅检索时，先抄一份只含目标 sheet 的干净工作簿，在副本上跑。

**只抄值，不抄公式结果以外的东西。** 行序必须逐行保持 —— 那几个脚本按 xlsx 行号
定位（`序号` 在合并表里不唯一），行序一变，A 件的答案就会写进 B 件的行且不报错。
故这里按 `iter_rows` 顺序逐行写，不排序、不去空行，写完断言行列数与源一致。
"""
from __future__ import annotations

import argparse
import pathlib
import sys

import openpyxl


def main() -> None:
    ap = argparse.ArgumentParser()
    ap.add_argument("--src", required=True)
    ap.add_argument("--sheet", required=True)
    ap.add_argument("--out", required=True)
    ap.add_argument("--force", action="store_true", help="允许覆盖已存在的 --out")
    args = ap.parse_args()

    src = pathlib.Path(args.src)
    out = pathlib.Path(args.out)
    if not src.exists():
        sys.exit(f"源文件不存在：{src}")
    if out.exists() and not args.force:
        sys.exit(f"目标已存在：{out}\n要覆盖请加 --force（覆盖前请确认它不是别人正在用的副本）")

    wb_in = openpyxl.load_workbook(src, read_only=True, data_only=True)
    if args.sheet not in wb_in.sheetnames:
        sys.exit(f"源文件里没有 sheet「{args.sheet}」。有的是：{wb_in.sheetnames}")
    ws_in = wb_in[args.sheet]

    wb_out = openpyxl.Workbook()
    ws_out = wb_out.active
    ws_out.title = args.sheet

    n_rows = 0
    n_cols = 0
    for row in ws_in.iter_rows(values_only=True):
        ws_out.append(list(row))
        n_rows += 1
        n_cols = max(n_cols, len(row))

    ws_out.freeze_panes = "A2"

    # 断言：行列数必须与源一致。少一行都可能让后续按行号定位的脚本写错行。
    if (ws_in.max_row, ws_in.max_column) != (n_rows, n_cols):
        sys.exit(f"抄写后维度对不上：源 {ws_in.max_row}×{ws_in.max_column}，"
                 f"新 {n_rows}×{n_cols} —— 已中止，未写出文件。")

    out.parent.mkdir(parents=True, exist_ok=True)
    wb_out.save(out)
    print(f"已写出 {out}")
    print(f"  sheet「{args.sheet}」 {n_rows} 行 × {n_cols} 列（含表头，行序与源逐行一致）")


if __name__ == "__main__":
    main()
