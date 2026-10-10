"""按 V3.0 第九、十节生成游览线路，中英各出一份 Excel。

    python3 route_plan.py --museum wmhg              # 写 exports/zh-CN/ 与 exports/en/
    python3 route_plan.py --museum wmhg --dry-run    # 只打印路线，不写文件

评级、Core、在展状态、名称从库里读；参观顺序、坐标、停留时间从 `<馆>_route_data.py` 读。
读库只读，不写任何表。

**算法（第九、十节）：**

    VisitScore   = G × clamp(Core + Ma + Mc + Mr, 0, 10)
    RouteUtility = VisitScore × AudienceFit ÷ (停留时间 + 新增步行时间)

- G：库里 `on_view` 为「未在展」时取 0，其余取 1。「未知」也取 1，导出时标出现状说明。
- Ma / Mc / Mr 一律取 0，AudienceFit 取 1：目前只做 General Visitor，季节、天气、修缮、
  拥挤都没有数据。同类重复（Mr 的一部分）也没做 —— 阶段一的同类组是逐批命名的，
  同一类东西散在几个组里（AGENTS.md 第 17 条已知问题 ①），拿它判重复不可靠。
- 选取：**先按评级**，S 全部考虑完才轮到 A，依此类推 —— Tier 本身就是参观优先级（第六节）。
  同一档之内有两种规则，每档时间两条都出（用户 2026-10-09 定）：
  · `utility`：第十节原公式，每次挑「边际效用」最高、放得进预算的一站
    （新增步行时间按加入后的整条路线重算）；
  · `score`：每次挑 VisitScore 最高、放得进预算的一站，不除以时间。
  两者的分歧在于 VisitScore 是按「一站」给的、与内容多少无关，除以时间之后
  45 分钟的大展天然吃亏：伪满皇宫的《从皇帝到公民》按 utility 要到全日才排得进。
- 顺序：选中的站一律按数据模块里的参观顺序走（第十节约束 3），从入口出发、回到入口。
- 步行：两点直线距离 × DETOUR ÷ WALK_M_PER_MIN。直线距离来自 OSM 坐标，
  绕行系数与步速都是估计。

数据对不上一律报错退出，不猜：库里的节点没在数据模块里登记、数据模块里的序号库里没有、
在展的藏品找不到所在的站，都直接 sys.exit。
"""

import argparse
import importlib
import math
import os
import re
import sys
from datetime import date

import openpyxl

import meta_lib
from export_excel import EN, ZH, mapv, safe_name, write_blocks

WALK_M_PER_MIN = 60    # 馆内缓行约 3.6 km/h（估计）
DETOUR = 1.4           # 直线距离折算成实际步行路径（估计）
TIER_RANK = {"S": 0, "A": 1, "B": 2, "C": 3}
RULES = {
    "utility": ("按效用（第十节原公式：分 ÷ 时间）",
                "By utility (Section 10 formula: score ÷ time)"),
    "score": ("按分（同一评级内分高者先）",
              "By score (highest score first within a tier)"),
}
RULE_SHORT = {"utility": ("按效用", "by utility"), "score": ("按分", "by score")}

CJK = re.compile(r"[一-鿿]")


def L(zh, en, lang):
    return zh if lang == ZH else en


# ---------------------------------------------------------------- 取数
def load_items(conn, mk):
    """{seq: dict}，评级取 COALESCE(tier_override, tier)，并与 artwork.tier 对账。"""
    cur = conn.cursor()
    cur.execute("""
        SELECT a.source_seq, t.object_type, COALESCE(t.tier_override, t.tier), a.tier,
               t.core, a.on_view, a.gallery_id,
               zh.text, en.text, e.tier_review_flag, e.review_reason, e.review_reason_en
        FROM artwork a
        JOIN museum m ON m.id = a.museum_id
        LEFT JOIN artwork_tier_v3 t ON t.museum_key = m.key_name AND t.source_seq = a.source_seq
        LEFT JOIN artwork_evidence e ON e.museum_key = m.key_name AND e.source_seq = a.source_seq
        LEFT JOIN content_text zh ON zh.content_id = a.name_cid AND zh.lang = 'zh-CN'
        LEFT JOIN content_text en ON en.content_id = a.name_cid AND en.lang = 'en'
        WHERE m.key_name = %s""", (mk,))
    items = {}
    for (seq, otype, tier, a_tier, core, on_view, gid,
         name_zh, name_en, flag, rr, rr_en) in cur.fetchall():
        if otype is None or core is None:
            sys.exit(f"seq {seq} 没有 V3 评分 —— 这个馆还没评完，不能排路线")
        if tier != a_tier:
            sys.exit(f"seq {seq} 评分表是 {tier}、artwork.tier 是 {a_tier}："
                     f"重灌后没跑 tier_v3_load.py --museum {mk} --apply-only")
        if not name_zh or not name_en:
            sys.exit(f"seq {seq} 缺{'中文' if not name_zh else '英文'}名称")
        items[seq] = dict(seq=seq, type=otype, tier=tier, core=float(core),
                          on_view=on_view, gallery_id=gid,
                          name={ZH: name_zh, EN: name_en},
                          review=bool(flag), review_reason={ZH: rr, EN: rr_en})
    if not items:
        sys.exit(f"库里没有馆 {mk}")
    return items


def check_registry(items, D):
    """库里每个节点都必须在数据模块里有去处；数据模块里的序号库里都得有。"""
    in_stops = [s[0] for s in D.STOPS]
    if len(in_stops) != len(set(in_stops)):
        sys.exit("STOPS 里有重复的序号")
    registered = set(in_stops) | set(D.MERGED) | set(D.UNPLACED)
    missing = set(registered) - set(items)
    if missing:
        sys.exit(f"数据模块登记了库里没有的序号：{sorted(missing)}")
    not_node = sorted(s for s in registered if items[s]["type"] != "node")
    if not_node:
        sys.exit(f"数据模块登记的这些序号在库里不是节点：{not_node}")
    nodes = {s for s, it in items.items() if it["type"] == "node"}
    orphan = nodes - registered
    if orphan:
        sys.exit("这些节点在数据模块里没有登记（STOPS / MERGED / UNPLACED 三选一）：\n  "
                 + "\n  ".join(f"{s} {items[s]['name'][ZH]}" for s in sorted(orphan)))
    for s, (into, _, _) in D.MERGED.items():
        if into not in in_stops:
            sys.exit(f"MERGED[{s}] 并入的 {into} 不在 STOPS 里")
    for s in in_stops:
        p = D.STOPS[in_stops.index(s)][1]
        if p not in D.POINTS:
            sys.exit(f"seq {s} 的坐标点 {p!r} 没有登记")


def highlights(items, D):
    """{站点 seq: [在展藏品]}：藏品所在展厅等于某站节点的展厅即归入该站。

    只认库里 on_view='在展' 且有 gallery_id 的藏品。归不进任何一站的直接报错 ——
    在展却排不进路线，说明数据模块漏了一站。
    """
    by_gallery = {}
    for s in (st[0] for st in D.STOPS):
        g = items[s]["gallery_id"]
        if g is not None:
            by_gallery.setdefault(g, s)       # 同一展厅有两站（怀远楼与清宴堂不同展厅，不会撞）
    out = {}
    for it in items.values():
        if it["type"] != "object" or it["on_view"] != "在展":
            continue
        s = by_gallery.get(it["gallery_id"])
        if s is None:
            sys.exit(f"藏品 seq {it['seq']} 在展，但所在展厅不对应任何一站")
        out.setdefault(s, []).append(it)
    for v in out.values():
        v.sort(key=lambda x: (TIER_RANK[x["tier"]], -x["core"], x["seq"]))
    return out


# ---------------------------------------------------------------- 算
def walk_min(a, b, D):
    la1, lo1 = D.POINTS[a][:2]
    la2, lo2 = D.POINTS[b][:2]
    r = 6371000.0
    p1, p2 = math.radians(la1), math.radians(la2)
    dp, dl = p2 - p1, math.radians(lo2 - lo1)
    h = math.sin(dp / 2) ** 2 + math.cos(p1) * math.cos(p2) * math.sin(dl / 2) ** 2
    meters = 2 * r * math.asin(math.sqrt(h))
    return meters * DETOUR / WALK_M_PER_MIN


class Planner:
    def __init__(self, items, D):
        self.D = D
        self.items = items
        self.order = [s[0] for s in D.STOPS]
        self.stop = {s[0]: s for s in D.STOPS}

    def visit_score(self, seq):
        it = self.items[seq]
        g = 0 if it["on_view"] == "未在展" else 1
        return g * max(0.0, min(10.0, it["core"]))   # Ma = Mc = Mr = 0

    def dwell(self, seq):
        return self.stop[seq][2]

    def legs(self, sel):
        """[(from_point, to_point, 分钟)]，入口出发、按参观顺序、回到入口。"""
        pts = [self.D.ENTRANCE] + [self.stop[s][1] for s in self.order if s in sel] \
            + [self.D.ENTRANCE]
        return [(a, b, walk_min(a, b, self.D)) for a, b in zip(pts, pts[1:])]

    def walk_total(self, sel):
        return sum(m for _, _, m in self.legs(sel))

    def total(self, sel):
        return sum(self.dwell(s) for s in sel) + self.walk_total(sel)

    def plan(self, budget, rule):
        """返回 (选中集合, {未选中的 seq: 原因 key})。rule 见 RULES。"""
        if rule not in RULES:
            sys.exit(f"未知的选站规则 {rule!r}")
        sel, why = set(), {}
        cands = [s for s in self.order if self.visit_score(s) > 0]
        for s in self.order:
            if self.visit_score(s) == 0:
                why[s] = "not_viewable"
        for tier in "SABC":
            pool = [s for s in cands if self.items[s]["tier"] == tier]
            while pool:
                base_walk = self.walk_total(sel)
                best = None
                for s in pool:
                    trial = sel | {s}
                    if self.total(trial) > budget:
                        continue
                    if rule == "utility":
                        added = self.walk_total(trial) - base_walk
                        u = self.visit_score(s) / (self.dwell(s) + added)
                    else:
                        u = self.visit_score(s)
                    key = (u, -self.order.index(s))   # 平手时先到的那站优先，结果确定
                    if best is None or key > best[0]:
                        best = (key, s)
                if best is None:
                    break
                sel.add(best[1])
                pool.remove(best[1])
            for s in pool:
                why[s] = "over_budget"
        return sel, why


# ---------------------------------------------------------------- 写表
TXT = {
    "file": ("路线_", "Routes - "),
    "sheet_about": ("说明", "About"),
    "sheet_all": ("候选总表", "All candidates"),
    "col_order": ("顺序", "#"),
    "col_rule": ("选站规则", "Selection rule"),
    "col_seq": ("序号", "Seq"),
    "col_name": ("站点", "Stop"),
    "col_tier": ("评级", "Tier"),
    "col_core": ("Core", "Core"),
    "col_vs": ("VisitScore", "VisitScore"),
    "col_dwell": ("建议停留（分钟）", "Suggested stay (min)"),
    "col_walk": ("自上一站步行（分钟）", "Walk from previous (min)"),
    "col_cum": ("累计（分钟）", "Cumulative (min)"),
    "col_status": ("现状", "Status"),
    "col_hl": ("在展藏品", "Objects on view here"),
    "col_dwell_basis": ("停留时间依据（估计）", "Basis for stay (estimate)"),
    "col_loc_src": ("地点出处", "Location source"),
    "col_coord": ("坐标", "Coordinates"),
    "col_review": ("需复核", "Needs review"),
    "col_type": ("类型", "Type"),
    "col_reason": ("未排入原因", "Why not included"),
    "entrance_start": ("出发", "Start"),
    "entrance_end": ("返回", "Return"),
    "sum_title": ("合计", "Totals"),
    "sum_budget": ("时间预算（分钟）", "Time budget (min)"),
    "sum_total": ("合计用时（分钟）", "Total time (min)"),
    "sum_dwell": ("其中停留", "of which visiting"),
    "sum_walk": ("其中步行", "of which walking"),
    "sum_stops": ("站数", "Stops"),
    "route_title": ("线路", "Route"),
    "node": ("节点", "Node"),
    "object": ("藏品", "Object"),
    "yes": ("是", "Yes"),
    "no": ("否", "No"),
    "in": ("✓", "✓"),
    "borrowed": ("坐标借用：", "Borrowed coordinates: "),
    "over_budget": ("超出时间预算", "Over time budget"),
    "not_viewable": ("不可参观", "Not viewable"),
    "obj_no_loc": ("官网未给出展出位置，在展状态未知", "No display location published; on-view status unknown"),
    "obj_in_stop": ("在「{}」站内展出，随该站参观，不单独计值",
                    "Shown inside '{}'; visited with that stop, not counted separately"),
    "key": ("项目", "Item"),
    "value": ("内容", "Detail"),
}


def T(k, lang):
    return TXT[k][0 if lang == ZH else 1]


def fmt(x):
    return round(x)


def hl_text(objs, lang):
    if not objs:
        return None
    return "；".join(f"{o['name'][lang]}（{o['tier']}）" for o in objs) if lang == ZH \
        else "; ".join(f"{o['name'][lang]} ({o['tier']})" for o in objs)


def status_text(P, s, lang):
    st = P.stop[s]
    if st[6]:
        return st[6] if lang == ZH else st[7]
    return mapv("on_view", P.items[s]["on_view"], lang)


def route_rows(P, sel, lang, hls):
    D = P.D
    rows, cum, prev = [], 0.0, D.ENTRANCE
    rows.append([T("entrance_start", lang), None, L(*D.ENTRANCE_NAME, lang)]
                + [None] * 11)
    for i, s in enumerate([x for x in P.order if x in sel], 1):
        st, it = P.stop[s], P.items[s]
        w = walk_min(prev, st[1], D)
        cum += w + st[2]
        p = D.POINTS[st[1]]
        coord = f"{p[0]:.5f},{p[1]:.5f} (OSM {p[2]})"
        if p[3]:
            coord = T("borrowed", lang) + L(p[3], p[4], lang) + " — " + coord
        src = D.SOURCES[st[5]]
        rows.append([
            i, s, it["name"][lang], it["tier"], round(it["core"], 2), round(P.visit_score(s), 2),
            st[2], fmt(w), fmt(cum), status_text(P, s, lang),
            hl_text(hls.get(s), lang),
            L(st[3], st[4], lang),
            L(src[0], src[1], lang) + (f" {src[2]}" if src[2] else ""),
            coord,
        ])
        prev = st[1]
    w = walk_min(prev, D.ENTRANCE, D)
    cum += w
    rows.append([T("entrance_end", lang), None, L(*D.ENTRANCE_NAME, lang),
                 None, None, None, None, fmt(w), fmt(cum)] + [None] * 5)
    return rows


ROUTE_COLS = ["col_order", "col_seq", "col_name", "col_tier", "col_core", "col_vs",
              "col_dwell", "col_walk", "col_cum", "col_status", "col_hl",
              "col_dwell_basis", "col_loc_src", "col_coord"]


def cols(keys):
    return [TXT[k] for k in keys]


def about_rows(P, mk_name, lang):
    D = P.D
    srcs = "；".join(f"{v[0]} {v[2]}" for v in D.SOURCES.values() if v[2]) if lang == ZH \
        else "; ".join(f"{v[1]} {v[2]}" for v in D.SOURCES.values() if v[2])
    return [
        [L("博物馆", "Museum", lang), mk_name],
        [L("生成日期", "Generated", lang), date.today().isoformat()],
        [L("开放时间", "Opening hours", lang), L(*D.OPENING, lang)],
        [L("观众", "Audience", lang), "General Visitor"],
        [L("公式", "Formula", lang),
         "VisitScore = G × clamp(Core + Ma + Mc + Mr, 0, 10); "
         "RouteUtility = VisitScore × AudienceFit ÷ (stay + added walking)"],
        [L("取值", "Parameters", lang),
         L("G：「未在展」为 0，其余为 1；Ma、Mc、Mr 均为 0，AudienceFit 为 1（尚无观众、季节、拥挤数据）",
           "G = 0 if not on view, else 1; Ma, Mc, Mr = 0 and AudienceFit = 1 "
           "(no audience, seasonal or crowding data yet)", lang)],
        [L("选取规则", "Selection", lang),
         L("先按评级（S → A → B → C）；同一评级内每档时间出两条线路：「按效用」每次挑"
           "分 ÷（停留 + 新增步行）最高的一站，「按分」每次挑分最高的一站，都要放得进预算。"
           "选中的站按馆方推荐的参观顺序排列，从入口出发、回到入口",
           "By tier first (S → A → B → C). Within a tier, each budget gets two routes: "
           "'by utility' repeatedly adds the stop with the highest score ÷ (stay + added walking), "
           "'by score' adds the highest-scoring stop; both only add stops that still fit. "
           "Selected stops follow the museum's recommended order, starting and ending at the "
           "entrance", lang)],
        [L("两种规则的差别", "Why two rules", lang),
         L("分数按「一站」给，与内容多少无关；除以时间后，45 分钟的大展排在 5 分钟的小景点之后",
           "Scores are per stop regardless of how much there is to see, so dividing by time "
           "puts a 45-minute exhibition behind a 5-minute stop", lang)],
        [L("预算口径", "Budget covers", lang),
         L("只算停留与步行，不含排队、安检、用餐、休息",
           "Visiting and walking only; excludes queues, security, meals and breaks", lang)],
        [L("步行时间", "Walking time", lang),
         L(f"OSM 坐标直线距离 × {DETOUR} ÷ 每分钟 {WALK_M_PER_MIN} 米（系数与步速为估计）",
           f"Straight-line distance between OSM coordinates × {DETOUR} ÷ {WALK_M_PER_MIN} m/min "
           "(factor and speed are estimates)", lang)],
        [L("停留时间", "Stay times", lang),
         L("全部为估计，依据官网写明的面积、展线长度与展品数量，无实测数据",
           "All estimated from published floor areas, display lengths and object counts; "
           "no measured data", lang)],
        [L("未做的约束", "Not yet handled", lang),
         L("同类重复（同类组划分不可靠）、预约、无障碍、单向通行、临时闭馆",
           "Repetition of similar stops (peer groups are unreliable), booking, accessibility, "
           "one-way routes, temporary closures", lang)],
        [L("出处", "Sources", lang), srcs],
    ]


def all_rows(P, plans, hls, lang):
    D, rows = P.D, []
    in_stop = {o["seq"]: s for s, objs in hls.items() for o in objs}
    for it in sorted(P.items.values(),
                     key=lambda x: (x["type"] != "node", TIER_RANK[x["tier"]], -x["core"], x["seq"])):
        s = it["seq"]
        marks, reason = [], None
        if s in P.stop:
            for key, *_ in D.BUDGETS:
                for rule in RULES:
                    sel, why = plans[key, rule]
                    marks.append(T("in", lang) if s in sel else T(why[s], lang))
            dwell = P.dwell(s)
        else:
            marks = [None] * (len(D.BUDGETS) * len(RULES))
            dwell = None
            if s in D.MERGED:
                reason = L(D.MERGED[s][1], D.MERGED[s][2], lang)
            elif s in D.UNPLACED:
                reason = L(*D.UNPLACED[s], lang)
            elif s in in_stop:
                reason = T("obj_in_stop", lang).format(P.items[in_stop[s]]["name"][lang])
            else:
                reason = T("obj_no_loc", lang)
        rr = it["review_reason"][lang] if it["review"] else None
        status = status_text(P, s, lang) if s in P.stop else mapv("on_view", it["on_view"], lang)
        rows.append([s, it["name"][lang], T(it["type"], lang), it["tier"], round(it["core"], 2),
                     status, dwell, *marks, reason,
                     T("yes", lang) if it["review"] else T("no", lang), rr])
    return rows


def write(P, plans, hls, mk_name, lang, out_root):
    D = P.D
    wb = openpyxl.Workbook()
    wb.remove(wb.active)
    write_blocks(wb.create_sheet(T("sheet_about", lang)),
                 [(None, cols(["key", "value"]), about_rows(P, mk_name, lang))], lang)
    for key, budget, lab_zh, lab_en in D.BUDGETS:
        summary, blocks = [], []
        for rule, rlab in RULES.items():
            sel, _ = plans[key, rule]
            dwell = sum(P.dwell(s) for s in sel)
            walk = P.walk_total(sel)
            summary.append([L(*rlab, lang), budget, fmt(dwell + walk), dwell, fmt(walk), len(sel)])
            blocks.append((T("route_title", lang) + L("：", ": ", lang) + L(*rlab, lang),
                           cols(ROUTE_COLS), route_rows(P, sel, lang, hls)))
        write_blocks(wb.create_sheet(L(lab_zh, lab_en, lang)), [
            (T("sum_title", lang),
             cols(["col_rule", "sum_budget", "sum_total", "sum_dwell", "sum_walk", "sum_stops"]),
             summary)] + blocks, lang)
    budget_cols = [(f"{b[2]}·{RULE_SHORT[r][0]}", f"{b[3]} · {RULE_SHORT[r][1]}")
                   for b in D.BUDGETS for r in RULES]
    all_cols = cols(["col_seq", "col_name", "col_type", "col_tier", "col_core", "col_status",
                     "col_dwell"]) + budget_cols + cols(["col_reason", "col_review"]) \
        + [("复核原因", "Review reason")]
    write_blocks(wb.create_sheet(T("sheet_all", lang)),
                 [(None, all_cols, all_rows(P, plans, hls, lang))], lang)
    d = os.path.join(out_root, lang)
    os.makedirs(d, exist_ok=True)
    path = os.path.join(d, safe_name(T("file", lang) + mk_name) + ".xlsx")
    wb.save(path)
    return path


def scan(path):
    wb = openpyxl.load_workbook(path, read_only=True)
    leaks = []
    for ws in wb.worksheets:
        for i, row in enumerate(ws.iter_rows(values_only=True), 1):
            for v in row:
                if isinstance(v, str) and CJK.search(v):
                    leaks.append(f"[{ws.title}] 行{i}: {v[:80]}")
    wb.close()
    return leaks


# ---------------------------------------------------------------- main
def main():
    ap = argparse.ArgumentParser(description=__doc__.split("\n\n")[0])
    ap.add_argument("--museum", required=True, help="museum.key_name，须有 <key>_route_data.py")
    ap.add_argument("--out", default="exports")
    ap.add_argument("--dry-run", action="store_true", help="只打印路线，不写文件")
    args = ap.parse_args()

    try:
        D = importlib.import_module(f"{args.museum}_route_data")
    except ModuleNotFoundError:
        sys.exit(f"没有 {args.museum}_route_data.py —— 这个馆还没有路线数据")

    conn = meta_lib.connect()
    items = load_items(conn, args.museum)
    cur = conn.cursor()
    cur.execute("SELECT zh.text, en.text FROM museum m"
                " JOIN content_text zh ON zh.content_id = m.name_cid AND zh.lang = 'zh-CN'"
                " JOIN content_text en ON en.content_id = m.name_cid AND en.lang = 'en'"
                " WHERE m.key_name = %s", (args.museum,))
    mk_zh, mk_en = cur.fetchone()
    conn.close()

    check_registry(items, D)
    hls = highlights(items, D)
    P = Planner(items, D)
    plans = {(key, rule): P.plan(budget, rule)
             for key, budget, *_ in D.BUDGETS for rule in RULES}

    for (key, rule), (sel, why) in plans.items():
        budget, lab = next((b[1], b[2]) for b in D.BUDGETS if b[0] == key)
        walk = P.walk_total(sel)
        dwell = sum(P.dwell(s) for s in sel)
        print(f"\n== {lab}·{rule}（预算 {budget}）：{len(sel)} 站，停留 {dwell} + 步行 {walk:.0f}"
              f" = {dwell + walk:.0f} 分钟")
        print("   " + " → ".join(f"{items[s]['name'][ZH]}[{items[s]['tier']}]"
                                  for s in P.order if s in sel))
        left = [s for s in P.order if s not in sel]
        if left:
            print("   未排入：" + "、".join(f"{items[s]['name'][ZH]}[{items[s]['tier']}]"
                                         for s in left))
    borrowed = [s for s in P.order if D.POINTS[P.stop[s][1]][3]]
    print(f"\n坐标借用 {len(borrowed)} 站：" + "、".join(items[s]["name"][ZH] for s in borrowed))
    print("不进路线的节点：" + "、".join(f"{items[s]['name'][ZH]}（{'并入' if s in D.MERGED else '位置不明'}）"
                                   for s in sorted(set(D.MERGED) | set(D.UNPLACED))))
    if args.dry_run:
        return

    for lang, name in ((ZH, mk_zh), (EN, mk_en)):
        path = write(P, plans, hls, name, lang, args.out)
        print(f"-> {path}")
    leaks = scan(path)                       # 最后写的是英文版
    if leaks:
        print(f"\n⚠ 英文版残留中文 {len(leaks)} 处：")
        for x in leaks[:30]:
            print("   " + x)
        sys.exit(4)
    print("英文版零中文残留")


if __name__ == "__main__":
    main()
