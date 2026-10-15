#!/usr/bin/env python3
"""导出宛平城本地资料、V3 评分与独立审计；不连接数据库。"""
import hashlib
import json
from collections import Counter
from pathlib import Path

from openpyxl import Workbook, load_workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

HERE=Path(__file__).resolve().parent
def load(name):return json.loads((HERE/name).read_text('utf-8'))

def main():
    source=load('wanping_site_data.json')
    sha=hashlib.sha256((HERE/'wanping_site_data.json').read_bytes()).hexdigest()
    rating_file='wanping_grouped_ratings.json' if (HERE/'wanping_grouped_ratings.json').exists() else 'wanping_ratings.json'
    rating=load(rating_file);am=load('wanping_audit_manifest.json');translated=load('wanping_translations.json')
    if any(v['input_sha256']!=sha for v in [rating,am,translated]):raise ValueError('输入版本不一致')
    if am.get('rating_file')!=rating_file:raise ValueError('审计对应的评分版本与导出版本不一致')
    if am.get('rating_sha256') and am['rating_sha256']!=hashlib.sha256((HERE/rating_file).read_bytes()).hexdigest():
        raise ValueError('审计对应的评分内容已经变化')
    if rating_file=='wanping_grouped_ratings.json':
        policy=load('wanping_peer_groups.json')
        if policy['state']!='用户已确认' or rating['grouping_policy_sha256']!=hashlib.sha256((HERE/'wanping_peer_groups.json').read_bytes()).hexdigest():
            raise ValueError('分组与用户确认版本不一致')
    ratings={r['seq']:r for r in rating['results']}
    raw_audit=list(map(json.loads,(HERE/'wanping_audit1_slim.jsonl').read_text('utf-8').splitlines()))
    audit={r['seq']:r for r in raw_audit}
    want={r['source_seq'] for r in source['rows']}
    if len(rating['results'])!=len(want) or len(raw_audit)!=len(want) or set(ratings)!=want or set(audit)!=want:raise ValueError('评分/审计漏行、多行或重复序号')
    tr=translated['values']
    def en(text):
        if not text:return ''
        return tr[text]['en']
    images={r['identity_key']:r for r in load('wanping_images.json')['images']}
    names={r['identity_key']:r['name'] for r in source['rows']}
    wb=Workbook();wb.remove(wb.active)
    headers_zh=['序号','关联整体','参观点名称','对象类型','Tier','Core','HS','IU','VI','VA','CE','ER','CR','需复核','审计原因与资料缺口','官方原文摘录','编者说明','图片链接','来源页面','评分模型','审计模型','同类组','评分依据','CR理由','S-ness理由']
    headers_en=['Seq','Related whole / area','Place / object','Object type','Tier','Core','HS','IU','VI','VA','CE','ER','CR','Needs review','Audit notes and missing evidence','Introduction (AI translation)','Editorial notes (AI translation)','Image URL','Source pages','Scoring model','Auditor','Peer group code']
    groups={g:f'PG{i:02d}' for i,g in enumerate(sorted({v['peer_group'] for v in ratings.values()}),1)}
    final=[]
    for lang,headers in [('zh-CN',headers_zh),('en',headers_en)]:
        ws=wb.create_sheet(lang);ws.append(headers)
        for r in source['rows']:
            s=r['source_seq'];v=ratings[s];a=audit[s]
            flag=a['found_factual_error'] or any(x['tier_sensitive'] for x in a['missing'])
            if lang=='zh-CN':
                notes='\n'.join(([a.get('error_note_zh')] if a.get('error_note_zh') else [])+[x['zh'] for x in a['missing']])
                name,parent,desc,note=r['name'],names.get(r['parent_key'],''),r['description'],r['editor_note']
            else:
                required=[(a.get('error_note_zh'),a.get('error_note_en'))]+[(x['zh'],x.get('en')) for x in a['missing']]
                if any(zh and not eng for zh,eng in required):raise ValueError(f'seq{s}:审计缺英译')
                notes='\n'.join(([a.get('error_note_en')] if a.get('error_note_en') else [])+[x['en'] for x in a['missing']])
                name,parent,desc,note=en(r['name']),en(names.get(r['parent_key'],'')),en(r['description']),en(r['editor_note'])
            image=images.get(r['identity_key'],{})
            image_url=image.get('url','') if image.get('retrieval')=='web_open_success' else ''
            values=[s,parent,name,v['object_type'],v['tier'],round(v['core'],4),v['HS'],v['IU'],v['VI'],v['VA'],v['CE'],v['ER'],round(v['CR'],4),('是' if flag else '否') if lang=='zh-CN' else ('Yes' if flag else 'No'),notes,desc,note,image_url,'\n'.join(dict.fromkeys(e['url'] for e in r['evidence'])),rating['model'],am['model'],v['peer_group'] if lang=='zh-CN' else groups[v['peer_group']]]
            if lang=='zh-CN':values += [v['evidence'],v['cr_reason'],v['sness_reason']]
            ws.append(values)
            if lang=='zh-CN':final.append(dict(r,rating=v,audit=a,needs_review=flag,image_reference=image or None))
        format_sheet(ws)
    refs=wb.create_sheet('来源证据');refs.append(['序号','名称','来源方式','URL','原文'])
    for r in source['rows']:
        for e in r['evidence']:refs.append([r['source_seq'],r['name'],e['retrieval_method'],e['url'],e['quote']])
    format_sheet(refs)
    gs=wb.create_sheet('同类组');gs.append(['组代码','同类组名','条目序号','参观点','状态'])
    for group,code in groups.items():
        members=[r for r in source['rows'] if ratings[r['source_seq']]['peer_group']==group]
        status='用户2026-10-15确认归并口径；其余组沿用原组' if rating_file=='wanping_grouped_ratings.json' else '原模型分组；展览整体与部分的归并待选择'
        gs.append([code,group,','.join(str(r['source_seq']) for r in members),' / '.join(r['name'] for r in members),status])
    format_sheet(gs);gs.column_dimensions['D'].width=110
    info=wb.create_sheet('说明');info.append(['项目','内容'])
    notes=[('范围',source['scope']),('采集日期',source['collected_on']),('状态','本地资料、评分及审计；未导入生产业务表，未上线导览站。'),
           ('来源','10份已取得正文的官方网页及1份官方索引摘录；原文与编者说明分列。'),
           ('对象关系','整体与组成部分有包含关系；不是39个互不重叠的路线站点。'),
           ('开放状态','不把历史报道及长期展览安排当作实时开放保证；未知项保留未知。'),
           ('英文出处','所有英文名称、简介、编者说明和审计自由文本均为AI翻译，不标成人工校对或馆方英文。'),
           ('图片','3个整体的官方资料图链接已读取；未下载或核实拍摄日期。其余条目无已核实的单独图片。'),
           ('CR分组',rating.get('grouping_note','模型原组名；已记录分批起组的局限。')),
           ('模型',rating['model']+' / '+am['model'])]
    for pair in notes:info.append(pair)
    format_sheet(info);info.column_dimensions['A'].width=22;info.column_dimensions['B'].width=120
    out=HERE/'exports'/'展品_宛平城及周边_本地审计版.xlsx'
    wb.save(out)
    final_path=HERE/'wanping_final.json'
    final_path.write_text(json.dumps(dict(input_sha256=sha,rating_file=rating_file,scope=source['scope'],state='本地评分与独立审计完成；未入生产业务表',models={'rating':rating['model'],'audit':am['model']},rows=final),ensure_ascii=False,indent=2)+'\n',encoding='utf-8')
    check=load_workbook(out,read_only=True)
    assert check['zh-CN'].max_row-1==check['en'].max_row-1==39
    for lang in ['zh-CN','en']:
        assert [r[0] for r in check[lang].iter_rows(min_row=2,values_only=True)]==sorted(want)
    residual=[]
    for row in check['en'].iter_rows(min_row=2,values_only=True):
        for i in [1,2,14,15,16]:
            if any('\u4e00'<=c<='\u9fff' for c in str(row[i] or '')):residual.append([row[0],i+1])
    summary={'records':39,'tier':dict(Counter(r['rating']['tier'] for r in final)),
             'types':dict(Counter(r['rating']['object_type'] for r in final)),
             'needs_review':[r['source_seq'] for r in final if r['needs_review']],
             'factual_errors':[r['source_seq'] for r in final if r['audit']['found_factual_error']],
             'english_cjk_residual':residual,'output':str(out)}
    (HERE/'wanping_export_report.json').write_text(json.dumps(summary,ensure_ascii=False,indent=2)+'\n',encoding='utf-8')
    print(json.dumps(summary,ensure_ascii=False))

def format_sheet(ws):
    ws.freeze_panes='D2';ws.auto_filter.ref=ws.dimensions
    for i in range(1,ws.max_column+1):ws.column_dimensions[get_column_letter(i)].width=16
    for i in [2,3]:ws.column_dimensions[get_column_letter(i)].width=34
    for i in [15,16,17,19]:
        if i<=ws.max_column:ws.column_dimensions[get_column_letter(i)].width=85
    for row in ws:
        for c in row:c.font=Font(name='Arial',size=11);c.alignment=Alignment(vertical='top',wrap_text=True)
    for c in ws[1]:c.font=Font(name='Arial',bold=True,color='FFFFFF');c.fill=PatternFill('solid',fgColor='234E52')
    for n in range(2,ws.max_row+1):ws.row_dimensions[n].height=100
    if ws.title=='来源证据':ws.column_dimensions['D'].width=65;ws.column_dimensions['E'].width=110

if __name__=='__main__':main()
